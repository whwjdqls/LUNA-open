"""Independently audit diagnostic scores and package the portable report on CPU."""

import argparse
import csv
import json
import os
import socket
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlsplit

import numpy as np
from audit_neuman_final_report import Links, digest, slide_previews
from PIL import Image
from pptx import Presentation

WORK = Path("/scratch2/whwjdqls99/LUNA-open")
BASE = WORK / "baselines/lhm-20260928"
KEYS = ["psnr", "l1", "foreground_l1", "background_l1", "ssim", "lpips", "mask_iou"]
RAW = {
    "lhm": BASE / "lhm-500m-test",
    "lhmpp": WORK / "baselines/lhmpp-20260928/lhmpp-700m-test",
    "ours": BASE / "ours-identity-14750-test",
}


def read(path):
    return json.loads(path.read_text())


def main():
    if not os.environ.get("SLURM_JOB_ID") or "login" in socket.gethostname():
        raise RuntimeError("Slurm compute node required")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    out = args.output
    protocol = read(out / "evidence/protocol.json")
    expected = {(s, f) for s, p in protocol["scenes"].items() for f in p["targets"]}
    csv_rows = list(csv.DictReader((out / "quantitative/metrics.csv").open()))
    frame_rows = list(csv.DictReader((out / "quantitative/per-frame.csv").open()))
    variants = sorted({r["variant"] for r in csv_rows})
    assert len(csv_rows) == len(variants) * 7
    assert len(frame_rows) == len(variants) * 41
    checks = {}
    for name in variants:
        d = read(out / "evidence" / f"{name}.json")
        assert len(d["frames"]) == 41
        assert {(r["scene"], r["frame"]) for r in d["frames"]} == expected
        for scene, values in {**d["per_scene"], "macro": d["mean_over_scenes"]}.items():
            row = next(r for r in csv_rows if r["variant"] == name and r["scene"] == scene)
            for k in KEYS:
                records = (
                    list(d["per_scene"].values())
                    if scene == "macro"
                    else [r["metrics"] for r in d["frames"] if r["scene"] == scene]
                )
                assert abs(np.mean([r[k] for r in records]) - values[k]) < 1e-10
                assert abs(float(row[k]) - values[k]) < 1e-10
        max_error = {k: 0.0 for k in ("l1", "psnr", "foreground_l1", "background_l1", "mask_iou")}
        source = RAW[name.removeprefix("raw-")] if name.startswith("raw-") else Path(d["root"])
        for r in d["frames"]:
            scene, frame = r["scene"], r["frame"]
            if name.startswith("canvas-"):
                gt_folder = WORK / "diagnostics/canvas-20260928/gt" / scene
                gt_path, mask_path = gt_folder / frame, gt_folder / f"mask-{frame}"
            else:
                gt_folder = BASE / "protocol-test" / scene
                gt_path, mask_path = gt_folder / "rgb" / frame, gt_folder / "mask" / frame
            with Image.open(gt_path) as im:
                gt = np.asarray(im.convert("RGB"), dtype=np.float64) / 255
            with Image.open(source / scene / "rgb" / frame) as im:
                pred = np.asarray(im.convert("RGB"), dtype=np.float64) / 255
            with Image.open(mask_path) as im:
                mask = np.asarray(im.convert("L"), dtype=np.float64) / 255
            with Image.open(source / scene / "alpha" / frame) as im:
                alpha = np.asarray(im.convert("L"), dtype=np.float64) / 255
            error = np.abs(pred - gt)
            values = dict(
                l1=float(error.mean()),
                psnr=float(-10 * np.log10(max(np.square(error).mean(), 1e-10))),
                foreground_l1=float((error * mask[..., None]).sum() / max(mask.sum() * 3, 1)),
                background_l1=float(
                    (error * (1 - mask[..., None])).sum() / max((1 - mask).sum() * 3, 1)
                ),
                mask_iou=float(
                    ((alpha >= 0.5) & (mask >= 0.5)).sum()
                    / max(((alpha >= 0.5) | (mask >= 0.5)).sum(), 1)
                ),
            )
            csv_row = next(
                x
                for x in frame_rows
                if x["variant"] == name and x["scene"] == scene and x["frame"] == frame
            )
            for k in KEYS:
                assert abs(float(csv_row[k]) - r["metrics"][k]) < 1e-10
            for k, value in values.items():
                diff = abs(value - r["metrics"][k])
                max_error[k] = max(max_error[k], diff)
                assert diff < (1e-4 if k == "psnr" else 2e-6), (name, scene, frame, k, diff)
        checks[name] = max_error
    print(
        f"Independent RGB/mask and aggregation checks passed for {len(variants) * 41} records",
        flush=True,
    )
    origins = read(out / "evidence/body-frames.json")
    for scene in protocol["scenes"]:
        v = origins["scenes"][scene]
        assert (
            np.max(
                np.abs(
                    np.array(v["smpl_pelvis"])
                    - v["smplx_pelvis"]
                    - np.array(v["canonical_root_translation"])
                )
            )
            < 1e-7
        )
        for m in RAW:
            assert len(list((out / "qualitative" / scene / "canonical" / m).glob("*.png"))) == 12
    for method in ("lhm", "lhmpp"):
        native = read(out / "evidence" / f"native-{method}-audit.json")
        assert len(native["native_fk_vs_fit_body"]) == 47
        assert max(r["max_mm"] for r in native["native_fk_vs_fit_body"]) < 0.0005
        pose = read(out / "evidence" / f"pose-oracle-{method}-audit.json")
        assert len(pose["records"]) == 41 and pose["fit_uses_test_rgb"]
        for r in pose["records"]:
            assert r["stages"]["articulated"]["mse"] <= r["stages"]["root"]["mse"] + 1e-7
            assert r["stages"]["root"]["mse"] <= r["initial_mse"] + 1e-7
            for s in r["stages"].values():
                assert len(s["mse_curve"]) == pose["steps"] + 1
                assert abs(min(s["mse_curve"]) - s["mse"]) < 1e-9
    image_count = 0
    for p in out.rglob("*.png"):
        with Image.open(p) as im:
            im.load()
        image_count += 1
    gallery = read(out / "evidence/gallery.json")
    assert len(gallery) == 42
    for item in gallery:
        with Image.open(out / item["path"]) as im:
            assert im.n_frames == item["frames"], item
            for n in range(im.n_frames):
                im.seek(n)
                im.load()
    prs = Presentation(out / "Alignment-investigation.pptx")
    for i, s in enumerate(prs.slides):
        for shape in s.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 100
            assert shape.top + shape.height <= prs.slide_height + 100
    slide_previews(prs, out)
    (out / "NeuMan-final-report-preview.pdf").rename(out / "Alignment-investigation-preview.pdf")
    audit = dict(
        host=socket.gethostname(),
        job=os.environ["SLURM_JOB_ID"],
        variants=len(variants),
        frames_per_variant=41,
        metrics=checks,
        canonical_views=216,
        decoded_pngs=image_count,
        gifs=len(gallery),
        slides=len(prs.slides),
        pptx_layout="bounds and Pillow text/media preview passed; Office playback not executed",
        status="passed",
    )
    (out / "AUDIT.json").write_text(json.dumps(audit, indent=2))
    links = Links()
    links.feed((out / "REPORT.html").read_text())
    local_links = 0
    for value in links.links:
        parsed = urlsplit(value)
        if parsed.scheme or parsed.netloc:
            continue
        assert (out / unquote(parsed.path)).exists(), value
        local_links += 1
    audit["local_links"] = local_links
    (out / "AUDIT.json").write_text(json.dumps(audit, indent=2))
    hashes = {
        str(p.relative_to(out)): digest(p)
        for p in sorted(out.rglob("*"))
        if p.is_file() and p.name != "SHA256.json"
    }
    (out / "SHA256.json").write_text(json.dumps(hashes, indent=2))
    archive = out.with_suffix(".zip")
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=5) as z:
        for path in sorted(out.rglob("*")):
            if path.is_file():
                z.write(path, path.relative_to(out.parent))
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
    receipt = dict(
        path=str(archive),
        bytes=archive.stat().st_size,
        sha256=digest(archive),
        crc="passed",
        files=len(hashes) + 1,
        audit=audit,
    )
    archive.with_suffix(".receipt.json").write_text(json.dumps(receipt, indent=2))
    print(json.dumps({k: v for k, v in receipt.items() if k != "audit"}, indent=2), flush=True)


if __name__ == "__main__":
    main()
