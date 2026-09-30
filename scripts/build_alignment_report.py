"""Package the alignment investigation, recalculated metrics, GIFs and PPT addendum.

Run on an allocated CPU node after all diagnostic variants have been scored.
The frozen benchmark is preserved. Test-RGB oracle fits are explicitly labeled.
"""

import argparse
import csv
import json
import os
import shutil
import socket
from pathlib import Path

import markdown2
from add_alignment_pose_table import add_pose_comparison_slide
from build_lhm_comparison_report import panel
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

WORK = Path("/scratch2/whwjdqls99/LUNA-open")
BASE = WORK / "baselines/lhm-20260928"
PLUS = WORK / "baselines/lhmpp-20260928"
DIAG = WORK / "diagnostics"
KEYS = ["psnr", "l1", "foreground_l1", "background_l1", "ssim", "lpips", "mask_iou"]
RAW = {
    "lhm": BASE / "lhm-500m-test",
    "lhmpp": PLUS / "lhmpp-700m-test",
    "ours": BASE / "ours-identity-14750-test",
}
LABELS = {"lhm": "LHM", "lhmpp": "LHM++", "ours": "Our identity 14750"}


def read(path):
    return json.loads(path.read_text())


def rgb(path):
    with Image.open(path) as image:
        return image.convert("RGB")


def csv_write(path, rows):
    with path.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main():
    if not os.environ.get("SLURM_JOB_ID") or "login" in socket.gethostname():
        raise RuntimeError("Slurm compute node required")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    out = args.output
    out.mkdir(parents=True, exist_ok=False)
    for folder in ("qualitative", "quantitative", "evidence/scripts", "slide-previews"):
        (out / folder).mkdir(parents=True)
    protocol = read(BASE / "protocol-test/protocol.json")
    scenes = list(protocol["scenes"])
    expected = {(s, f) for s, p in protocol["scenes"].items() for f in p["targets"]}
    assert len(expected) == 41
    sources = {"raw-" + m: p.with_name(p.name + "-metrics.json") for m, p in RAW.items()}
    for p in sorted((DIAG / "alignment-metrics-20260928").glob("*.json")):
        sources[p.stem] = p
    required = {f"pose-{m}-{v}" for m in ("lhm", "lhmpp") for v in ("root", "articulated")}
    assert required <= set(sources), required - set(sources)
    data, rows, frames = {}, [], []
    for name, source in sources.items():
        result = data[name] = read(source)
        assert len(result["frames"]) == 41
        assert {(r["scene"], r["frame"]) for r in result["frames"]} == expected
        oracle = bool(result.get("test_rgb_optimized", False))
        for scene, metrics in {**result["per_scene"], "macro": result["mean_over_scenes"]}.items():
            rows.append(
                dict(
                    variant=name,
                    scene=scene,
                    test_rgb_fitted=oracle,
                    **{k: metrics[k] for k in KEYS},
                )
            )
        for r in result["frames"]:
            frames.append(
                dict(
                    variant=name,
                    scene=r["scene"],
                    frame=r["frame"],
                    test_rgb_fitted=oracle,
                    **{k: r["metrics"][k] for k in KEYS},
                )
            )
        shutil.copy2(source, out / "evidence" / f"{name}.json")
    csv_write(out / "quantitative/metrics.csv", rows)
    csv_write(out / "quantitative/per-frame.csv", frames)
    for source, name in (
        (BASE / "protocol-test/protocol.json", "protocol.json"),
        (BASE / "fit-test/summary.json", "smplx-conversion-summary.json"),
        (BASE / "ours-identity-14750-test/canonical-camera.json", "canonical-camera.json"),
        (BASE / "identity-snapshots/identity-014750.json", "identity-snapshot.json"),
        (DIAG / "body-frames-20260928.json", "body-frames.json"),
        (DIAG / "alignment-20260928/alignment.json", "alignment-parameters.json"),
        (DIAG / "canvas-20260928/protocol.json", "canvas-protocol.json"),
        (DIAG / "canonical-ours-20260928/method.json", "canonical-ours-method.json"),
    ):
        shutil.copy2(source, out / "evidence" / name)
    for method in ("lhm", "lhmpp"):
        for category in ("native", "pose-oracle"):
            for filename in ("audit.json", "model-load.json"):
                source = DIAG / f"{category}-{method}-20260928" / filename
                shutil.copy2(source, out / "evidence" / f"{category}-{method}-{filename}")
    script_names = [
        "diagnose_neuman_alignment.py",
        "audit_neuman_body_frames.py",
        "diagnose_native_avatar.py",
        "diagnose_native_avatar_yonsei.sh",
        "prepare_alignment_canvas.py",
        "score_alignment_diagnostics.py",
        "score_alignment_yonsei.sh",
        "render_diagnostic_identity_orbit.py",
        "diagnose_lhm_pose_oracle.py",
        "build_alignment_report.py",
        "add_alignment_pose_table.py",
        "audit_alignment_report.py",
        "audit_neuman_final_report.py",
        "evaluate_lhm_neuman.py",
        "evaluate_lhmpp_neuman.py",
        "build_lhm_comparison_report.py",
    ]
    for name in script_names:
        shutil.copy2(Path(__file__).parent / name, out / "evidence/scripts" / name)
    (out / "evidence/logs").mkdir()
    for name in (
        "alignment-20260928.log",
        "body-frames-20260928.log",
        "native-lhm-diagnostic.log",
        "native-lhmpp-diagnostic.log",
        "alignment-metrics.log",
        "alignment-canvas.log",
        "pose-oracle-lhm.log",
        "pose-oracle-lhmpp.log",
        "pose-metrics-lhm.log",
        "pose-metrics-lhmpp.log",
    ):
        source = WORK / "logs" / name
        if source.exists():
            shutil.copy2(source, out / "evidence/logs" / name)
    gallery = []

    def save_gif(images, dest):
        images[0].save(dest, save_all=True, append_images=images[1:], duration=350, loop=0)
        gallery.append(dict(path=str(dest.relative_to(out)), frames=len(images)))

    for scene, info in protocol["scenes"].items():
        folder = out / "qualitative" / scene
        folder.mkdir()
        canonical = {
            m: DIAG
            / ("canonical-ours-20260928" if m == "ours" else f"native-{m}-20260928/canonical")
            / scene
            for m in RAW
        }
        orbit = []
        for angle in range(0, 360, 30):
            images = []
            for m in RAW:
                source = canonical[m] / f"{angle:03}.png"
                dest = folder / "canonical" / m
                dest.mkdir(parents=True, exist_ok=True)
                shutil.copy2(source, dest / source.name)
                images.append(rgb(source))
            orbit.append(
                panel(
                    images,
                    list(LABELS.values()),
                    f"{scene} | canonical T pose | yaw {angle:03} deg",
                    "Shared camera and pelvis; no canonical ground truth. Identity and shape remain native.",
                    tile=384,
                )
            )
        save_gif(orbit, folder / "canonical-aligned.gif")
        orbit[0].save(folder / "canonical-aligned.png")
        old_images = []
        for m in ("lhm", "lhmpp"):
            old_images += [rgb(canonical[m] / "unaligned.png"), rgb(canonical[m] / "000.png")]
        old_images.append(rgb(canonical["ours"] / "000.png"))
        before = panel(
            old_images,
            ["LHM before", "LHM aligned", "LHM++ before", "LHM++ aligned", "Our identity"],
            f"{scene} | canonical display correction",
            "SMPL-X pelvis translated into the SMPL frame (~12.7 cm up). This does not change posed benchmark scores.",
            tile=320,
        )
        before.save(folder / "canonical-before-after.png")
        for m in ("lhm", "lhmpp", "ours"):
            alignment = []
            for frame in info["targets"]:
                gt = rgb(BASE / "protocol-test" / scene / "rgb" / frame)
                images = [gt, rgb(RAW[m] / scene / "rgb" / frame)]
                images += [
                    rgb(DIAG / "alignment-20260928" / f"{m}-{v}" / scene / "rgb" / frame)
                    for v in ("translation", "similarity", "affine")
                ]
                alignment.append(
                    panel(
                        images,
                        [
                            "Ground truth",
                            "Original",
                            "2D translation*",
                            "2D similarity*",
                            "2D affine*",
                        ],
                        f"{scene}/{frame} | {LABELS[m]}",
                        "*Diagnostic oracle: fitted to test RGB. Changes image placement, not avatar appearance.",
                        tile=320,
                    )
                )
            save_gif(alignment, folder / f"alignment-{m}.gif")
            alignment[len(alignment) // 2].save(folder / f"alignment-{m}.png")
        for m in ("lhm", "lhmpp"):
            poses = []
            root = DIAG / f"pose-oracle-{m}-20260928"
            for frame in info["targets"]:
                original = root / "repeat" / scene / "rgb" / frame
                if not original.exists():
                    original = RAW[m] / scene / "rgb" / frame
                images = [rgb(BASE / "protocol-test" / scene / "rgb" / frame), rgb(original)]
                images += [rgb(root / v / scene / "rgb" / frame) for v in ("root", "articulated")]
                poses.append(
                    panel(
                        images,
                        ["Ground truth", "Original pose", "Root pose fit*", "Joint pose fit*"],
                        f"{scene}/{frame} | {LABELS[m]}",
                        "*Test-RGB oracle, frozen identity. Camera, shape, Gaussian attributes and weights unchanged.",
                        tile=384,
                    )
                )
                for v in ("root", "articulated"):
                    for kind in ("rgb", "alpha"):
                        dest = folder / "pose-renders" / m / v / kind
                        dest.mkdir(parents=True, exist_ok=True)
                        shutil.copy2(root / v / scene / kind / frame, dest / frame)
            save_gif(poses, folder / f"pose-{m}.gif")
            poses[len(poses) // 2].save(folder / f"pose-{m}.png")
        comparisons = []
        for frame in info["targets"]:
            gt = rgb(BASE / "protocol-test" / scene / "rgb" / frame)
            images = [gt] + [rgb(RAW[m] / scene / "rgb" / frame) for m in RAW]
            comparisons.append(
                panel(
                    images,
                    ["Ground truth"] + list(LABELS.values()),
                    f"{scene}/{frame} | frozen original benchmark",
                    "Common 512 crop; supplied NeuMan poses. No test-RGB pose or image alignment.",
                    tile=384,
                )
            )
            dest = folder / "ground-truth"
            dest.mkdir(exist_ok=True)
            gt.save(dest / frame)
        save_gif(comparisons, folder / "original-benchmark.gif")
        print(f"Packaged {scene}", flush=True)
    (out / "evidence/gallery.json").write_text(json.dumps(gallery, indent=2))

    def table(names):
        text = "| Variant | PSNR ↑ | L1 ↓ | LPIPS ↓ | FG L1 ↓ | IoU ↑ |\n|---|---:|---:|---:|---:|---:|\n"
        for key, label in names:
            d = data[key]["mean_over_scenes"]
            text += f"| {label} | {d['psnr']:.4f} | {d['l1']:.6f} | {d['lpips']:.6f} | {d['foreground_l1']:.6f} | {d['mask_iou']:.6f} |\n"
        return text

    base_table = table([(f"raw-{m}", LABELS[m]) for m in RAW])
    pose_table = table(
        [("pose-lhmpp-repeat", "LHM++ — same-reconstruction control (no RGB fit)")]
        + [
            (f"pose-{m}-{v}", f"{LABELS[m]} — {v} pose oracle*")
            for m in ("lhm", "lhmpp")
            for v in ("root", "articulated")
        ]
    )
    image_table = table(
        [
            (f"alignment-{m}-{v}", f"{LABELS[m]} — {v} oracle*")
            for m in RAW
            for v in ("translation", "similarity", "affine")
        ]
    )
    canvas_table = table(
        [
            (
                f"canvas-{m}-{v}",
                f"{LABELS[m]} — original canvas, {v}" + ("*" if v == "similarity" else ""),
            )
            for m in RAW
            for v in ("raw", "similarity")
        ]
    )
    native_table = table(
        [
            (f"native-{m}-{v}", f"{LABELS[m]} — {v}")
            for m in ("lhm", "lhmpp")
            for v in ("repeat", "zero-shape", "fixed-shape", "recrop")
        ]
    )
    lhm_pose = data["pose-lhm-articulated"]["mean_over_scenes"]
    pp_pose = data["pose-lhmpp-articulated"]["mean_over_scenes"]
    report = f"""# NeuMan alignment investigation and recalculated metrics

Yonsei • RTX 4090 job 2346749, node32 • CPU job 2346750, cnode02 • 28–29 September 2026 KST

## Findings

1. **The user-identified canonical misalignment is real.** Our shared camera was applied to different SMPL and SMPL-X template origins. SMPL-X's pelvis was about **12.6–12.7 cm lower**, approximately 26 pixels in the front display. The new canonical views translate the SMPL-X pelvis into the SMPL frame, using body-model joints only. No image fitting, height scaling, or appearance changes were used.
2. **That display bug did not remove translation from the scored poses.** Those renders already used the fitted SMPL-X translations. Across all 47 converted poses, native joint transforms agree with the standard SMPL-X forward model within **0.000427 mm**. The earlier SMPL-to-SMPL-X conversion averaged 3.52 mm in 3D and 0.71 pixels in projection. These checks establish faithful conversion of the supplied poses; they do not establish perfect alignment with the photographed person or reconstructed clothing.
3. **Pose compatibility materially affects the scores.** Keeping reconstructed identity, shape and camera fixed, fitting root and joint angles to each test image raises LHM to **{lhm_pose["psnr"]:.4f} PSNR**, and LHM++ to **{pp_pose["psnr"]:.4f}**. These are diagnostic oracle results, because they use held-out RGB. They cannot replace feed-forward test scores.
4. **The evaluation canvas also has a large effect.** Placing the unchanged LHM crop prediction back into the original NeuMan white image raises PSNR from 19.8613 to 23.1590 and changes LPIPS from 0.094337 to 0.054540. Pixel area, resampling and clipping all contribute. LUNA's exact crop, resolution, background handling, LPIPS backbone and score aggregation are not established by the available paper text.
5. **The exact Table 1 gap is not fully attributable yet.** Alignment and protocol choices explain substantial sensitivity. The released checkpoint variant, reference preprocessing, driving poses and evaluation code must all match before claiming reproduction. Better-looking surfaces need not have lower pixel errors when silhouette, texture placement and pose differ.

## 1. Correct target: LUNA Table 1

The target is reconstruction/self-reenactment, not the cross-identity experiment. [LUNA Table 1](https://arxiv.org/html/2606.31981v2#S4.T1) reports:

| Paper method | NeuMan PSNR ↑ | L1 ↓ | LPIPS ↓ |
|---|---:|---:|---:|
| LHM | 25.310 | 0.029 | 0.039 |
| MV-LHM* | 26.832 | 0.017 | 0.023 |
| LUNA | 26.819 | 0.015 | 0.023 |

**LHM++ is not MV-LHM.** It has no corresponding row in that table. Our model is the locally trained identity encoder with SMPL/LBS, not the complete LUNA animator. Our identity has trained on the six NeuMan identities; released baselines receive no local fine-tuning. Input counts remain LHM 1, LHM++ 4, ours 4.

LUNA §4.1 says one or four uniformly selected training references and official test splits. Figure 5 discusses SAM-3D-Body converted to SMPL(X) for its motion comparison; this does not fully specify the Table 1 fitting/evaluation pipeline. Here we use the supplied NeuMan SMPL annotations converted with the official correspondence assets. We cannot assume these are identical protocols.

## 2. Fixed benchmark and recalculation rules

All experiments use **41 official test frames in six scenes** (bike 10, citron 3, jogging 10, lab 10, parkinglot 4, seattle 4), and the same frozen training references. Our checkpoint stays **14,750** so changing checkpoint quality does not confound this investigation. The full dataset split is 344 train / 44 validation / 41 test.

Scores are computed from lossless 8-bit PNGs on [0,1], averaged per frame, then per scene, then equally over six scenes. PSNR uses RGB MSE; SSIM uses an 11×11 window; LPIPS is AlexNet on [-1,1]. NeuMan masks define GT/reference white compositing and crops. Predictions are not cleaned by GT masks. FG/BG L1 selects error regions only. LHM++ mask comes from DPT; other masks are rasterizer alpha.

The confirmed canonical display fix leaves the **posed benchmark numbers unchanged**. These remain the honest fixed-protocol scores:

{base_table}

Our original LHM L1 is already below the paper's 0.029 even though PSNR and LPIPS are worse. The discrepancy is not uniformly worse on every metric and should not be reduced to visual quality alone.

## 3. Canonical comparison, now in the same frame

The baseline native reconstruction uses its A-pose prior and native inverse bind; the display renders native zero/T pose. We preserve each method's learned Gaussians and fitted shape, align pelvis origins, and orbit the same camera through 12 angles. Differences in body shape are preserved. **NeuMan has no canonical T-pose ground-truth image/scan for these subjects**, so canonical PSNR/L1 would be undefined. Our render is not ground truth.

The old downloadable report is preserved as a historical record. Its canonical origin mismatch is corrected by this package. See the before/after and aligned orbit GIFs below.

## 4. Image alignment diagnostic

Each frame is independently fitted to test RGB using bounded MSE optimization: translation ±64 pixels; similarity adds scale 0.7–1.3 and rotation ±15 degrees; affine adds an independent y scale and shear ±0.15. White padding, fixed GT and no GT-mask cleanup. The known-translation numerical check recovered (11,-7) pixels within 0.00001 pixels.

{image_table}

*Every aligned result is a **test-RGB oracle**, not an official test score. Best full-resolution MSE candidates include the unchanged render, but this is a local bounded optimization, not a proven global optimum. LPIPS is not optimized and can become worse after interpolation. Approximately 66% of baseline MSE lies in a ±5-pixel GT silhouette boundary band (mean of 41 frame fractions, not the six-scene macro). This supports sensitivity to silhouette/pose/shape mismatch; it does not prove all error is camera translation.

## 5. 3D pose diagnostic with frozen avatars

Native rendering is differentiated with respect to six root parameters and then 21×3 joint parameters. Root axis-angle components are bounded to ±20 degrees, translation components to ±0.15 m, and body axis-angle components to ±25 degrees; these are additive axis-angle coordinates. Adam uses learning rate 0.035, 80 steps per stage; the best MSE is retained. Cameras, shape betas, identity network weights, Gaussian colors/scales/opacities and canonical geometry remain fixed. Native skinning changes posed geometry. LHM++'s fixed DPT can change view-dependent output as posed features change.

{pose_table}

*These fit the evaluated RGB and are **oracle diagnostics**. They test how much mismatch a different pose could explain. They do not demonstrate that an RGB-independent pose estimator would achieve these numbers, nor establish pose as the sole source of the paper gap. Per-frame parameters, optimization curves and gradients are recorded in evidence. LHM's initial control is the frozen benchmark with a negligible native repeat shift; LHM++ saves a control from the exact same reconstruction used by its pose fit. These are not converged upper bounds: 39/41 LHM articulated fits have their best MSE in the final ten iterations. Further optimization could explain more error; the remaining gap does not prove a residual appearance deficit. Median LHM translation correction after joint fitting is 7.50 cm.

Joint fitting can also compensate for reconstructed shape, clothing placement and occlusion differences. Its gains do not uniquely diagnose incorrect NeuMan joint annotations.

## 6. Native-canvas diagnostic

The recorded crop transform places each prediction back into the original NeuMan image dimensions using bilinear resize and white canvas. GT is the original NeuMan frame composited white with its supplied mask. Predictions receive no GT-mask cleanup. This tests an alternative scoring domain and resampling; it is not a discovered LUNA protocol or an accuracy improvement.

{canvas_table}

Similarity rows additionally use test-RGB alignment. The LHM original-canvas + similarity PSNR 24.4477 is closer to 25.310, while L1/LPIPS still differ. Choosing a canvas merely to match the paper would not validate the experiment.

## 7. Shape and reference input checks

{native_table}

Zero shape and first-reference fixed shape do not resolve the gap. Tall-reference recropping uses the reference GT mask bounding box plus a declared 5% margin at 840×504; target geometry/cameras stay unchanged. It does not help LHM, but improves LHM++ in this run. This is an observed preprocessing ablation, not proof of the exact optimal upstream pipeline.

Repeat controls differ slightly from the frozen renders: LHM macro PSNR +0.000413 dB; LHM++ +0.011918 dB. LHM++'s native image-token merging has its own advancing CUDA generator, and original/recrop runs are interleaved; those inputs are not paired by generator state. The crop effect is larger than the aggregate repeat shift, but a strict paired-RNG study would be needed for an isolated causal claim. Canonical outputs are fresh reconstructions and are not claimed bitwise identical to the earlier exports.

## 8. All qualitative results in one folder

Each row includes every test frame for that scene. GIFs cycle sparse official test frames; they are not densely sampled motion videos. Orbit GIFs contain 12 camera angles. PNG panels use a fixed middle frame, without picking by score.

| Scene | Aligned canonical | Canonical before/after | LHM pose oracle | LHM++ pose oracle | Original benchmark |
|---|---|---|---|---|---|
"""
    for s in scenes:
        report += f"| {s} | [GIF](qualitative/{s}/canonical-aligned.gif) | [PNG](qualitative/{s}/canonical-before-after.png) | [GIF](qualitative/{s}/pose-lhm.gif) | [GIF](qualitative/{s}/pose-lhmpp.gif) | [GIF](qualitative/{s}/original-benchmark.gif) |\n"
    for s in scenes:
        report += f"\n### {s}\n\n![{s} aligned canonical](qualitative/{s}/canonical-aligned.gif)\n\n[2D alignment: LHM](qualitative/{s}/alignment-lhm.gif) · [LHM++](qualitative/{s}/alignment-lhmpp.gif) · [Our identity](qualitative/{s}/alignment-ours.gif)\n"
    report += """
## 9. Quantitative files and presentation

- [All seven metrics by variant and scene, CSV](quantitative/metrics.csv)
- [All per-frame metrics, CSV](quantitative/per-frame.csv)
- [Editable PowerPoint addendum](Alignment-investigation.pptx)
- [Slide preview PDF](Alignment-investigation-preview.pdf)
- [Numerical and package audit](AUDIT.json)
- [Gallery frame-count manifest](evidence/gallery.json)
- [Body origins](evidence/body-frames.json), [conversion summary](evidence/smplx-conversion-summary.json), [camera](evidence/canonical-camera.json), [reference/target protocol](evidence/protocol.json)
- [LHM pose parameters](evidence/pose-oracle-lhm-audit.json), [LHM++ pose parameters](evidence/pose-oracle-lhmpp-audit.json), [2D transforms](evidence/alignment-parameters.json), [canvas mapping](evidence/canvas-protocol.json)

Raw pose-corrected RGB/alpha and all aligned canonical angle PNGs are copied into `qualitative/<scene>/`. The original diagnostic render directories remain under `/scratch2/whwjdqls99/LUNA-open/diagnostics/`; the metric JSONs record their paths.

## 10. Provenance and reproducibility

Source hierarchy: LUNA §3.3 and §4 motivate consulting its LHM baseline; LHM canonical reconstruction and LBS apply to the native released-model adapter. LHM++ is evaluated separately. Source pins: LHM `4f88aaeb3629249fbbddb4d0784a06962d9e1338`; LHM++ `906b5d9fb967ab42efb92f6fa55bf22cac86b653`. The adapters call upstream code and preserve their source/license files. Model weights, body assets and dataset terms are separate. No model/body weights are included in this report.

Execution uses Yonsei compute nodes only: CPU acquisition/analysis/packaging on cnode02, native PyTorch 2.3/cu118 and modern identity/metrics environments on one RTX4090 node32. Native code uses CUDA architecture 8.9. Our canonical renderer's first JIT build auto-selected the visible 8.9 device; its wrapper now makes that setting explicit. Identity training on node31 was not modified.

Evidence scripts are the final reference versions. Some scripts were formatted after execution; alignment JSON retains its executed pre-format SHA. The pose script was extended for LHM++ after the LHM process loaded it. These reference copies are not represented as byte-identical to every earlier execution. Metric JSONs, logs and per-frame outputs are the execution evidence.

Primary sources: [LUNA](https://arxiv.org/html/2606.31981v2), [LHM](https://arxiv.org/html/2503.10625v1), [LHM++](https://arxiv.org/html/2506.13766v2), [LHM source](https://github.com/aigc3d/LHM), [LHM++ source](https://github.com/aigc3d/LHM-plusplus).

### Next benchmark decision

Keep the frozen 512-crop benchmark as the current comparable result. Use the corrected canonical gallery for visual inspection. Report pose/image oracle and native-canvas scores in separate diagnostic tables. A paper-matched Table 1 reproduction still needs its exact reference selection, pose estimator/checkpoint variant and crop/scoring implementation; no test-RGB correction should be silently substituted into a feed-forward benchmark.
"""
    (out / "REPORT.md").write_text(report)
    html = markdown2.markdown(report, extras=["tables", "fenced-code-blocks"])
    (out / "REPORT.html").write_text(
        "<!doctype html><meta charset='utf-8'><title>NeuMan alignment investigation</title><style>body{max-width:1200px;margin:40px auto;padding:0 24px;font:17px/1.6 system-ui;color:#243448}img{max-width:100%}table{border-collapse:collapse;font-size:14px}td,th{padding:8px;border:1px solid #ced6de}a{color:#146ca4}h2{margin-top:48px}</style>"
        + html
    )

    # A short addendum to the previously delivered full progress presentation.
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def slide(title, lines=(), picture=None, note=""):
        s = prs.slides.add_slide(prs.slide_layouts[6])

        def text(x, y, w, h, content, size=20):
            box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            tf = box.text_frame
            tf.word_wrap = True
            for i, line in enumerate(content.split("\n")):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                p.font.size = Pt(size)
                p.font.color.rgb = RGBColor.from_string("243448")
                p.space_after = Pt(12)
            return box

        text(0.4, 0.15, 12.5, 0.8, title, 28)
        if picture:
            im = rgb(picture)
            factor = min(12.4 / im.width, 4.65 / im.height)
            w, h = im.width * factor, im.height * factor
            s.shapes.add_picture(
                str(picture),
                Inches((13.333 - w) / 2),
                Inches(1.1),
                width=Inches(w),
                height=Inches(h),
            )
            text(0.5, 6.1, 12.2, 0.95, "\n".join(lines), 17)
        else:
            text(0.65, 1.35, 12, 5.45, "\n".join(lines), 23)
        text(
            0.5,
            7.08,
            12,
            0.32,
            f"Yonsei / fixed 41 test frames / {len(prs.slides):02d} — {note}",
            10,
        )
        return s

    slide(
        "Why do attractive avatars score poorly?",
        [
            "Confirmed: canonical panels mixed SMPL and SMPL-X origins.",
            "Measured: pose compatibility and scoring canvas strongly affect the metrics.",
            "All scores recalculated on the same 41 official test frames.",
            "Oracle pose/image fits use test RGB; they are diagnostic only.",
        ],
    )
    slide(
        "LUNA Table 1 is the comparison target",
        [
            "LHM: PSNR 25.310 | L1 0.029 | LPIPS 0.039",
            "MV-LHM: 26.832 | 0.017 | 0.023 — not LHM++",
            "LUNA: 26.819 | 0.015 | 0.023",
            "Our identity + SMPL/LBS is not the complete LUNA model.",
            "Exact paper crop, scoring and fitting code remain unavailable.",
        ],
        note="Source: arxiv.org/html/2606.31981v2, Table 1",
    )
    slide(
        "Confirmed display bug: pelvis origins differ",
        [
            "SMPL-X pelvis is ~12.7 cm lower (~26 display pixels).",
            "The correction changes canonical display only; scored poses already contain fitted translation.",
        ],
        out / "qualitative/bike/canonical-before-after.png",
    )
    for s in scenes:
        slide(
            f"{s}: canonical avatars in a shared frame",
            [
                "Same camera, T pose and pelvis origin; native shape and appearance preserved.",
                f"12-angle animation: qualitative/{s}/canonical-aligned.gif. No canonical ground truth.",
            ],
            out / "qualitative" / s / "canonical-aligned.png",
        )
    slide(
        "Original benchmark remains the comparable score",
        [
            f"{LABELS[m]}: PSNR {data['raw-' + m]['mean_over_scenes']['psnr']:.3f} | L1 {data['raw-' + m]['mean_over_scenes']['l1']:.5f} | LPIPS {data['raw-' + m]['mean_over_scenes']['lpips']:.5f}"
            for m in RAW
        ]
        + [
            "512×512 crops, white background, six-scene equal mean.",
            "Baselines are released models; ours trained on the six identities.",
        ],
    )
    slide(
        "Pose oracle: same avatar, different fitted motion",
        [
            f"LHM joint fit: {lhm_pose['psnr']:.3f} PSNR | {lhm_pose['l1']:.5f} L1 | {lhm_pose['lpips']:.5f} LPIPS",
            "Fitting uses test RGB. Identity, camera and shape stay frozen.",
        ],
        out / "qualitative/bike/pose-lhm.png",
    )
    slide(
        "LHM++ pose sensitivity",
        [
            f"LHM++ joint fit: {pp_pose['psnr']:.3f} PSNR | {pp_pose['l1']:.5f} L1 | {pp_pose['lpips']:.5f} LPIPS",
            "Fitting uses test RGB. Frozen native neural renderer; no local fine-tuning.",
        ],
        out / "qualitative/bike/pose-lhmpp.png",
    )
    slide(
        "Canvas choice explains a large numerical shift",
        [
            "LHM crop: 19.861 PSNR / 0.02457 L1 / 0.09434 LPIPS",
            "Same render on original canvas: 23.159 / 0.01182 / 0.05454",
            "Original canvas + similarity oracle: 24.448 / 0.00988 / 0.05347",
            "This is protocol sensitivity, not a new avatar or proven LUNA protocol.",
            "See all seven metrics and scene/frame breakdowns in quantitative/.",
        ],
    )
    slide(
        "What is resolved, and what remains",
        [
            "Fixed: canonical origin mismatch; regenerated six complete orbit GIFs.",
            "Verified: native FK matches the converted SMPL-X body to <0.0005 mm.",
            "Measured: global image alignment, root/joint pose, shape and crop sensitivity.",
            "Unresolved: exact paper scoring domain, driving fits and released model variant.",
            "Use diagnostic tables separately; do not replace feed-forward scores with test fits.",
        ],
    )
    add_pose_comparison_slide(prs)
    prs.save(out / "Alignment-investigation.pptx")
    (out / "BUILD.json").write_text(
        json.dumps(
            dict(
                host=socket.gethostname(),
                job=os.environ["SLURM_JOB_ID"],
                variants=len(data),
                slides=len(prs.slides),
                gifs=len(gallery),
            ),
            indent=2,
        )
    )
    print(
        f"Built {out}: {len(data)} metric variants, {len(gallery)} GIFs, {len(prs.slides)} slides",
        flush=True,
    )


if __name__ == "__main__":
    main()
