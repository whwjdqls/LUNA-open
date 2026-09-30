"""Build the three-method NeuMan report, complete GIF gallery and editable PPTX."""

import argparse
import csv
import hashlib
import json
import os
import shutil
import socket
from pathlib import Path

import markdown2
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from build_lhm_comparison_report import panel
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

WORK = Path("/scratch2/whwjdqls99/LUNA-open")
REPO = Path("/home/whwjdqls99/LUNA-open")
BASE = WORK / "baselines/lhm-20260928"
PLUS = WORK / "baselines/lhmpp-20260928"
KEYS = ["psnr", "l1", "foreground_l1", "background_l1", "ssim", "lpips", "mask_iou"]
TITLES = ["PSNR (dB) ↑", "L1 ↓", "FG L1 ↓", "BG L1 ↓", "SSIM ↑", "LPIPS ↓", "IoU ↑"]
METHODS = [
    dict(
        id="lhm",
        label="LHM-500M",
        short="LHM",
        refs=1,
        color="#D58043",
        renders=BASE / "lhm-500m-test",
        metrics=BASE / "lhm-500m-test-metrics.json",
    ),
    dict(
        id="lhmpp",
        label="LHM++-700M + DPT",
        short="LHM++",
        refs=4,
        color="#268C77",
        renders=PLUS / "lhmpp-700m-test",
        metrics=PLUS / "lhmpp-700m-test-metrics.json",
    ),
    dict(
        id="ours",
        label="Our identity · 14750",
        short="Ours",
        refs=4,
        color="#356CAF",
        renders=BASE / "ours-identity-14750-test",
        metrics=BASE / "ours-identity-14750-test-metrics.json",
    ),
]


def read(path):
    return json.loads(path.read_text())


def gif(frames, path):
    frames[0].save(path, save_all=True, append_images=frames[1:], duration=300, loop=0)


def image(path):
    with Image.open(path) as im:
        return im.convert("RGB")


def write_csv(path, rows):
    with path.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main():
    if not os.environ.get("SLURM_JOB_ID") or "login" in socket.gethostname():
        raise RuntimeError("Slurm compute node required")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    out = args.output
    if out.exists():
        raise FileExistsError(out)
    for name in ("qualitative", "quantitative", "evidence/scripts", "slide-previews"):
        (out / name).mkdir(parents=True)
    protocol = read(BASE / "protocol-test/protocol.json")
    scenes = list(protocol["scenes"])
    expected = {(s, f) for s, info in protocol["scenes"].items() for f in info["targets"]}
    assert len(expected) == 41
    camera_hash = hashlib.sha256(
        (BASE / "ours-identity-14750-test/canonical-camera.json").read_bytes()
    ).hexdigest()
    metric_rows, frame_rows = [], []
    for method in METHODS:
        method["data"] = data = read(method["metrics"])
        assert data["split"] == "test"
        assert data["method"]["manifest_sha256"] == protocol["manifest_sha256"]
        if method["id"] != "ours":
            assert data["method"]["canonical_camera_sha256"] == camera_hash
        assert {(r["scene"], r["frame"]) for r in data["frames"]} == expected
        assert len(data["frames"]) == 41
        for scene, info in protocol["scenes"].items():
            assert data["method"]["reference_frames"][scene] == info["references"][: method["refs"]]
        for scene, metrics in {**data["per_scene"], "macro": data["mean_over_scenes"]}.items():
            metric_rows.append(
                dict(
                    method=method["label"],
                    references=method["refs"],
                    scene=scene,
                    **{key: metrics[key] for key in KEYS},
                )
            )
        for row in data["frames"]:
            frame_rows.append(
                dict(
                    method=method["label"],
                    scene=row["scene"],
                    frame=row["frame"],
                    **{key: row["metrics"][key] for key in KEYS},
                )
            )
        shutil.copy2(method["metrics"], out / "evidence" / f"{method['id']}-metrics.json")
        for name in ("method.json", "model-load.json", "projection-check.json"):
            if (method["renders"] / name).exists():
                shutil.copy2(method["renders"] / name, out / "evidence" / f"{method['id']}-{name}")
    write_csv(out / "quantitative/metrics.csv", metric_rows)
    write_csv(out / "quantitative/per-frame.csv", frame_rows)
    for src, name in (
        (BASE / "protocol-test/protocol.json", "protocol.json"),
        (BASE / "fit-test/fits.json", "conversion-fits.json"),
        (BASE / "ours-identity-14750-test/canonical-camera.json", "canonical-camera.json"),
        (BASE / "identity-snapshots/identity-014750.json", "identity-snapshot.json"),
        (PLUS / "requirements-resolved.txt", "lhmpp-environment.txt"),
        (PLUS / "asset-links.json", "lhmpp-assets.json"),
        (PLUS / "shared-body-asset-check.json", "lhmpp-shared-body-assets.json"),
        (WORK / "assets/lhmpp-priors/acquisition.json", "lhmpp-prior-acquisition.json"),
        (WORK / "assets/lhmpp_700m/config.json", "lhmpp-released-config.json"),
        (WORK / "assets/lhmpp_700m-receipt.json", "lhmpp-checkpoint-acquisition.json"),
        (WORK / "assets/lhmpp_700m/model-verified-sha256.json", "lhmpp-checkpoint-hash.json"),
    ):
        shutil.copy2(src, out / "evidence" / name)
    for name in (
        "evaluate_lhmpp_neuman.py",
        "run_lhmpp_neuman_yonsei.sh",
        "setup_lhmpp_yonsei.sh",
        "prepare_lhmpp_priors.py",
        "evaluate_lhm_neuman.py",
        "evaluate_renders.py",
        "fit_neuman_smplx.py",
        "build_neuman_final_report.py",
        "build_lhm_comparison_report.py",
        "audit_neuman_final_report.py",
        "yonsei_env.sh",
    ):
        shutil.copy2(REPO / "scripts" / name, out / "evidence/scripts" / name)
    source = WORK / "runs/neuman-identity-v2-20260928/source/src/luna_open/identity_training.py"
    shutil.copy2(source, out / "evidence/identity-training-source.py")

    # Every benchmark metric gets a per-scene chart; raw numbers remain editable CSV.
    for stem, selected, layout in (
        ("primary-metrics", [0, 1, 2, 5], (2, 2)),
        ("additional-metrics", [3, 4, 6], (1, 3)),
    ):
        fig, axs = plt.subplots(
            *layout, figsize=(14, 7 if layout[0] == 2 else 4), constrained_layout=True
        )
        for ax, metric_index in zip(np.asarray(axs).ravel(), selected):
            key = KEYS[metric_index]
            for i, method in enumerate(METHODS):
                ax.bar(
                    np.arange(len(scenes)) + (i - 1) * 0.25,
                    [method["data"]["per_scene"][s][key] for s in scenes],
                    0.24,
                    label=f"{method['short']} · {method['refs']} ref",
                    color=method["color"],
                )
            ax.set_title(TITLES[metric_index], fontsize=13)
            ax.set_xticks(np.arange(len(scenes)), scenes, rotation=30, ha="right")
            ax.grid(axis="y", alpha=0.18)
        np.asarray(axs).ravel()[0].legend(fontsize=8)
        for ext in ("png", "svg"):
            fig.savefig(out / f"quantitative/{stem}.{ext}", dpi=180)
        plt.close(fig)

    media = []
    for scene, info in protocol["scenes"].items():
        folder = out / "qualitative" / scene
        folder.mkdir()
        frames = []
        for method in METHODS:
            for kind in ("rgb", "alpha", "references"):
                shutil.copytree(method["renders"] / scene / kind, folder / method["id"] / kind)
            cano_path = (
                method["renders"] / scene / "qualitative/canonical-front.png"
                if method["id"] == "ours"
                else method["renders"] / scene / "canonical/front.png"
            )
            shutil.copy2(cano_path, folder / method["id"] / "canonical-front.png")
            if method["id"] != "ours":
                shutil.copytree(
                    method["renders"] / scene / "canonical",
                    folder / method["id"] / "canonical",
                )
        (folder / "ground-truth").mkdir()
        (folder / "ground-truth-mask").mkdir()
        for name in info["targets"]:
            gt = BASE / "protocol-test" / scene / "rgb" / name
            shutil.copy2(gt, folder / "ground-truth" / name)
            shutil.copy2(
                BASE / "protocol-test" / scene / "mask" / name,
                folder / "ground-truth-mask" / name,
            )
            comparison = panel(
                [image(gt)] + [image(m["renders"] / scene / "rgb" / name) for m in METHODS],
                ["Ground truth", "LHM · 1 ref", "LHM++ + DPT · 4 refs", "Our identity · 4 refs"],
                f"{scene} | {name} | Official NeuMan test",
                "Same target camera and body annotations; identity checkpoint 14750; 512 x 512 crops",
            )
            comparison.save(folder / name)
            frames.append(comparison)
        gif(frames, folder / "comparison.gif")
        canonical = panel(
            [image(folder / m["id"] / "canonical-front.png") for m in METHODS],
            ["LHM · 1 reference", "LHM++ + DPT · 4 refs", "Our identity · 4 refs"],
            f"{scene} | Canonical front views",
            "Shared virtual camera; no canonical ground truth; these views are not scored",
        )
        canonical.save(folder / "canonical.png")
        shutil.copy2(
            BASE / "ours-identity-14750-test" / scene / "qualitative/gt-canonical-lbs.gif",
            folder / "ours-gt-canonical-lbs.gif",
        )
        lhmpp_frames = [
            panel(
                [
                    image(folder / "ground-truth" / name),
                    image(folder / "lhmpp/canonical-front.png"),
                    image(folder / "lhmpp/rgb" / name),
                ],
                ["Ground truth", "LHM++ canonical", "LHM++ posed + DPT"],
                f"{scene} | {name}",
                "Canonical identity reused for all target poses; SMPL-X LBS + DPT",
            )
            for name in info["targets"]
        ]
        gif(lhmpp_frames, folder / "lhmpp-gt-canonical-posed.gif")
        references = panel(
            [image(BASE / "protocol-test" / scene / "rgb" / name) for name in info["references"]],
            info["references"],
            f"{scene} | Evaluation reference images",
            "LHM: first image. LHM++ and our identity: all four. All references belong to the training split.",
            tile=256,
        )
        references.save(folder / "references.png")
        middle = info["targets"][len(info["targets"]) // 2]
        media.append(
            dict(
                scene=scene,
                gif=f"qualitative/{scene}/comparison.gif",
                selected=f"qualitative/{scene}/{middle}",
                canonical=f"qualitative/{scene}/canonical.png",
                targets=info["targets"],
            )
        )
        gallery = "<!doctype html><meta charset='utf-8'><style>body{font:18px sans-serif;margin:24px}img{max-width:100%}</style>"
        gallery += f"<h1>{scene}: all {len(info['targets'])} official test frames</h1><a href='../../REPORT.html'>Report</a>"
        gallery += "<p>Ground truth | LHM | LHM++ with DPT | Our identity 14750</p><img src='comparison.gif'>"
        for name in info["targets"]:
            gallery += f"<h2>{name}</h2><img src='{name}'>"
        (folder / "index.html").write_text(gallery)
    for start in (0, 3):
        rows = [image(out / m["selected"]) for m in media[start : start + 3]]
        rows = [im.resize((1600, round(im.height * 1600 / im.width))) for im in rows]
        sheet = Image.new("RGB", (1600, sum(im.height for im in rows)), "white")
        y = 0
        for im in rows:
            sheet.paste(im, (0, y))
            y += im.height
        sheet.save(out / f"qualitative/overview-{start // 3 + 1}.png")

    metrics_table = (
        "| Method | Refs | " + " | ".join(TITLES) + " |\n|---|---:|" + "---:|" * len(KEYS) + "\n"
    )
    for method in METHODS:
        m = method["data"]["mean_over_scenes"]
        metrics_table += (
            "| "
            + method["label"]
            + f" | {method['refs']} | "
            + " | ".join(f"{m[k]:.4f}" if k == "psnr" else f"{m[k]:.6f}" for k in KEYS)
            + " |\n"
        )
    lhmpp = METHODS[1]["data"]["mean_over_scenes"]
    lhm = METHODS[0]["data"]["mean_over_scenes"]
    delta = f"LHM++ versus LHM: PSNR {lhmpp['psnr'] - lhm['psnr']:+.2f} dB; L1 {lhmpp['l1'] - lhm['l1']:+.5f}; LPIPS {lhmpp['lpips'] - lhm['lpips']:+.5f}."
    report = f"""# NeuMan final comparison: LHM, LHM++, and our identity encoder

[Editable PowerPoint](NeuMan-final-report.pptx) · [All metrics CSV](quantitative/metrics.csv) ·
[Per-frame CSV](quantitative/per-frame.csv) · [Slide outline](SLIDES.md)

## Measured results

All **41 official test frames**, six scenes, equal-weight mean over scene means.
Identity checkpoint **14750** is frozen from the earlier comparison; selection used validation LPIPS.
{delta}

{metrics_table}

![Primary metrics](quantitative/primary-metrics.png)
![Additional metrics](quantitative/additional-metrics.png)

## Step 1 — Fix the benchmark

NeuMan has 344 training, 44 validation and 41 test frames in this manifest.
Test counts: bike 10, citron 3, jogging 10, lab 10, parkinglot 4, seattle 4.
Every method uses identical target frames, 512×512 person crops, white backgrounds,
crop intrinsics and body annotations. NeuMan's provided foreground segmentations
define the input/target crops and white input backgrounds, including test cropping.
This is an oracle-mask preprocessing protocol. No target RGB is used to reconstruct identity.
The fixed reference filenames, annotations and manifest fingerprint are in [protocol.json](evidence/protocol.json).

## Step 2 — Preserve each method

- **LHM-500M:** released checkpoint; one reference; native Gaussian reconstruction and SMPL-X skinning; no local fine-tuning.
- **LHM++-700M:** separate released checkpoint; the same four references as ours; native point/image transformer, diffused voxel SMPL-X skinning and **DPT neural renderer**. This is not LUNA's MV-LHM and not the PixelShuffle variant.
- **Our identity:** four references; trained on the training frames of these six identities; SMPL body template and annotation-driven LBS teacher; the neural animator is outside this comparison.

LHM++ uses the common square reference images; its native DINO wrapper resizes 512 to 504 pixels.
This differs from the upstream tall reference-canvas defaults and is recorded as a benchmark adaptation.
LHM++'s DPT requires dimensions divisible by seven. Its native feature renderer runs at 518×518
with **unchanged** full crop intrinsics, then the extra bottom/right pixels are removed.
No target-mask cleanup, target-color fitting or alignment to test RGB is applied.
LHM++ directly predicts final RGB with a learned white background and a separate mask;
the RGB is exported directly, as upstream does. The predicted DPT mask supplies its IoU.
LHM and ours use their rasterized alpha. These mask outputs come from different stages.

## Step 3 — Match body poses and audit geometry

NeuMan SMPL annotations are converted to SMPL-X using the user's official transfer assets.
The 47 audited fits cover all test frames plus each scene's first reference.
Mean mesh correspondence residual is 3.52 mm; mean reprojection difference is 0.71 pixels.
Conversion optimizes body geometry only. Both released baselines use these same conversions.
Their native deformation implementations remain distinct.
[Conversion audit](evidence/conversion-fits.json) · [LHM++ camera check](evidence/lhmpp-projection-check.json).

## Step 4 — Score lossless predictions

- PSNR: per-frame −10 log10 of full-crop RGB MSE on [0,1], then scene means.
- L1: mean absolute RGB error over the full crop.
- FG/BG L1: the same error averaged over GT foreground/background pixels separately; predictions are unchanged.
- SSIM: 11×11 Gaussian window, sigma 1.5. LPIPS: AlexNet, RGB mapped to [−1,1].
- IoU: predicted and GT masks thresholded at 0.5.
- Average frames within each scene, then average six scenes equally. No reference frames are scored.

All seven previously reported metrics are included. Canonical renders have no ground truth and are unscored.
These sparse test-frame GIFs do not establish temporal stability; MAE/MSJ are outside this image benchmark.
[LHM records](evidence/lhm-metrics.json) · [LHM++ records](evidence/lhmpp-metrics.json) · [Our records](evidence/ours-metrics.json).

## Step 5 — Explain our training inputs

The current identity training sampler randomly chooses a scene and eight distinct training frames:
four references produce one canonical reconstruction, supervised through four other target frames.
Four groups give 16 targets per optimizer update. References and targets are disjoint within a group;
reference selection is resampled across groups. A shared network is trained across all six identities.
Cached frozen Sapiens features are per frame; caching does not fix the reference set.
Validation and test evaluation use four fixed training-split references per sequence.
Training losses include balanced foreground/background RGB L1, LPIPS, silhouette L1 and Gaussian priors.
[Frozen training source](evidence/identity-training-source.py) · [Exact configuration](evidence/ours-method.json).

## Interpretation and limitations

Metric rankings measure agreement with exact held-out frames. They are not a standalone realism ranking.
Inspect smoothness, face/clothing detail and articulation alongside numeric scores.
LHM++ improves full-crop L1, SSIM, LPIPS and background L1 over LHM, but has worse
PSNR, foreground L1 and predicted-mask IoU. Its foreground L1 is 0.1620 versus
LHM's 0.1301, while background L1 falls to 0.0030 from 0.0099. Full-crop L1 mixes
these two regions; a cleaner background can coexist with worse person reconstruction.
PSNR depends on squared error, so lower absolute error does not guarantee higher PSNR.
Our model has the best six other aggregate metrics; LHM++ has the lowest background L1.

Visual inspection of the fixed middle test frame in each scene shows smoother faces,
garments and silhouettes in LHM++ than our visibly rough Gaussian output.
Parkinglot's dark hair and white headphones are clearer than in LHM; jogging's red
jacket has more coherent folds. These improvements coexist with target mismatches:
the bike body is narrow and the head direction differs, lab's head is more upright
than GT, and seattle's head is more in profile. These are observations of the overview
frames, with every test frame included below for inspection. Our lower metric errors
do not remove its visible splat artifacts.
We do not attribute a numerical gap to a single cause without an ablation.
Our identity has local training exposure to all six people; the released baselines have no local fine-tuning.
Released pretraining overlap with NeuMan is unknown. LHM has one reference, LHM++ and ours have four.
This evaluates new frames of familiar identities for our model, not unseen-identity generalization.
It is a local identity-plus-posing benchmark, not a completed reproduction of LUNA's neural animation results.

## Qualitative results

Each comparison GIF shows **ground truth | LHM | LHM++ + DPT | our identity**.
All official test frames are included in chronological order; playback is 300 ms per frame for inspection,
not original video timing. Raw RGB, masks, references, canonical views and individual comparison PNGs
are included under each scene. GIFs are also embedded in the PowerPoint.
"""
    for m in media:
        s = m["scene"]
        report += f"\n### {s}\n\n![{s}]({m['gif']})\n\n[Every frame](qualitative/{s}/index.html) · [References](qualitative/{s}/references.png) · [Canonical comparison]({m['canonical']}) · [LHM++ GT/canonical/posed GIF](qualitative/{s}/lhmpp-gt-canonical-posed.gif) · [Our GT/canonical/LBS GIF](qualitative/{s}/ours-gt-canonical-lbs.gif)\n"
    report += """
## Sources and reproducibility

LUNA §3.3 motivates identity pretraining through an MV-LHM extension; this experiment keeps
released LHM++ separate. See [LUNA](https://arxiv.org/html/2606.31981v2),
[LHM](https://arxiv.org/html/2503.10625v1), [LHM++](https://arxiv.org/html/2506.13766v2).
LHM++ source: [commit 906b5d9](https://github.com/aigc3d/LHM-plusplus/tree/906b5d9fb967ab42efb92f6fa55bf22cac86b653).
Weights: [LHMPP-700M, fc9f7366](https://huggingface.co/3DAIGC/LHMPP-700M/tree/fc9f73664b9bfcc457e96210ddc06fe7caf0d559).
Native code called: ModelHumanA4OLRM.infer_single_view; GSPlatBackFeatRenderer.forward_animate_gs;
PatchDPT4DecoderOnly.forward. Exact loading, configuration, hashes, resource records and adapters are in evidence/.
Source, weights, body assets and NeuMan retain their separate licenses. No weights/body assets are included here.
The evaluations ran on Yonsei RTX 4090 allocations; no B200 measurements are transferred.
Runtime numbers have different scopes (our model consumes cached features), so they are not a speed benchmark.

PowerPoint GIF animation is intended for Slide Show mode; viewers can display a static first frame.
The separate GIFs and HTML gallery are the portable playback alternative. See PACKAGE_CHECKS.json for verification scope.
"""
    (out / "REPORT.md").write_text(report)
    css = "body{font:17px/1.55 system-ui;max-width:1440px;margin:36px auto;padding:0 24px;color:#243448}h1,h2,h3{color:#173450}img{max-width:100%}table{border-collapse:collapse;font-size:15px}td,th{padding:10px;border:1px solid #ccd4dc}th{background:#edf2f7}a{color:#246cb0}"
    (out / "REPORT.html").write_text(
        "<!doctype html><meta charset='utf-8'><title>NeuMan final comparison</title><style>"
        + css
        + "</style>"
        + markdown2.markdown(report, extras=["tables"])
    )
    make_slides(out, protocol, media, delta)
    (out / "README.txt").write_text(
        "Open NeuMan-final-report.pptx for the editable final deck, or REPORT.html for the complete local report. GIFs are embedded in the deck and also provided in qualitative/. Keep the folder together when downloading. All seven benchmark metrics are in quantitative/metrics.csv and per-frame.csv. Identity checkpoint 14750 is fixed. No model weights or licensed body assets are included.\n"
    )
    (out / "build-summary.json").write_text(
        json.dumps(
            dict(
                host=socket.gethostname(),
                job=os.environ["SLURM_JOB_ID"],
                methods=3,
                test_frames=41,
                scenes=scenes,
                gif_count=18,
            ),
            indent=2,
        )
    )
    print(out, flush=True)


def make_slides(out, protocol, media, delta):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    outline, specs = [], []

    def text(slide, x, y, w, h, content, size=22, bold=False, color="243448"):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        for i, line in enumerate(content.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "DejaVu Sans"
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(15)
        return box

    def slide(title, bullets=None, asset=None, caption="", table=None, notes=""):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        number = len(prs.slides)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        text(s, 0.5, 0.3, 12.3, 0.72, title, 28, True, "173450")
        if bullets:
            text(s, 0.7, 1.35, 11.9, 5.4, "\n".join("• " + b for b in bullets), 22)
        if asset:
            im = image(out / asset)
            scale = min(12.2 / im.width, 5.25 / im.height)
            w, h = im.width * scale, im.height * scale
            s.shapes.add_picture(
                str(out / asset),
                Inches((13.333 - w) / 2),
                Inches(1.12 + (5.25 - h) / 2),
                Inches(w),
                Inches(h),
            )
            text(s, 0.55, 6.45, 12.2, 0.6, caption, 14)
        if table:
            rows, widths = table
            shape = s.shapes.add_table(
                len(rows), len(rows[0]), Inches(0.6), Inches(1.65), Inches(12.1), Inches(2.5)
            )
            t = shape.table
            for c, width in enumerate(widths):
                t.columns[c].width = Inches(width)
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c)
                    cell.text = str(val)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(
                        "EAF0F6" if r == 0 else "F8FAFC"
                    )
                    for p in cell.text_frame.paragraphs:
                        p.font.name = "DejaVu Sans"
                        p.font.size = Pt(17)
                        p.font.bold = r == 0
                        p.font.color.rgb = RGBColor.from_string("173450")
            text(s, 0.65, 4.65, 12, 1.4, caption, 19)
        footer = text(
            s,
            0.55,
            7.15,
            11.8,
            0.2,
            f"Yonsei · NeuMan · 41 official test frames · identity step 14750                         {number:02d}",
            9,
            color="687B8F",
        )
        footer.click_action.hyperlink.address = "REPORT.html"
        s.notes_slide.notes_text_frame.text = (
            notes
            + "\n"
            + caption
            + ("\nAsset: " + asset if asset else "")
            + "\nFull protocol and evidence: REPORT.html"
        ).strip()
        specs.append(
            dict(
                number=number,
                title=title,
                bullets=bullets,
                asset=asset,
                caption=caption,
                table=table,
                notes=notes,
            )
        )
        outline.extend(
            [
                f"## {number}. {title}",
                "",
                *(bullets or []),
                "",
                caption,
                "",
                f"[Media]({asset})" if asset else "",
                "",
                notes,
                "",
            ]
        )

    slide(
        "NeuMan: LHM, LHM++ and our identity encoder",
        bullets=[
            "Final report of the local identity reconstruction comparison.",
            "Three methods · six people · all 41 official test frames.",
            "Seven metrics, complete comparison GIFs and canonical views.",
            "Our identity checkpoint is fixed at update 14,750.",
            "Posing uses body annotations; our neural animator is outside this benchmark.",
        ],
        notes="This report adds released LHM++ to the previously completed LHM/local-identity benchmark. It is not a reproduction of LUNA's complete animation results.",
    )
    slide(
        "What we completed",
        bullets=[
            "1. Fixed the official split, reference images, crop cameras and masks.",
            "2. Preserved native released LHM and LHM++ inference.",
            "3. Converted SMPL annotations to SMPL-X and audited the fits.",
            "4. Evaluated all methods with one common image metric implementation.",
            "5. Packaged every test frame, animated comparisons and reproducible evidence.",
        ],
    )
    slide(
        "Three different reconstruction systems",
        table=(
            [
                ["Method", "Refs", "Local training", "Posing / output"],
                ["LHM-500M", "1", "None", "SMPL-X / Gaussian RGB"],
                ["LHM++-700M", "4", "None", "SMPL-X / native DPT RGB"],
                ["Our identity 14750", "4", "Six NeuMan identities", "SMPL / Gaussian RGB"],
            ],
            [2.7, 0.7, 3.4, 5.3],
        ),
        caption="LHM++ is a separate released model. LUNA's MV-LHM is a distinct, unpublished extension.",
        notes="LUNA §3.3 and §4 motivate the identity reconstruction comparison. LHM++ source commit 906b5d9; checkpoint fc9f7366. Released pretraining overlap with NeuMan is unknown.",
    )
    slide(
        "The fixed benchmark",
        bullets=[
            "Official split: 344 train / 44 validation / 41 test frames.",
            "Test counts: bike 10; citron 3; jogging 10; lab 10; parkinglot 4; seattle 4.",
            "Same target frames, 512 × 512 crops, white background and crop intrinsics.",
            "Average frames within each scene, then weight all six scenes equally.",
            "Identity checkpoint selected on validation LPIPS; test frames did not select it.",
        ],
        notes="evidence/protocol.json and evidence/identity-snapshot.json",
    )
    slide(
        "Reference images: fixed only for evaluation",
        asset="qualitative/bike/references.png",
        caption="Bike: 00000, 00034, 00069, 00103. LHM uses the first; LHM++ and ours use all four.",
        notes="Every selected reference belongs to the official training split. The same identity reconstruction is reused for every target frame of a scene.",
    )
    slide(
        "How our identity model was trained",
        bullets=[
            "Randomly choose one of the six sequences and eight distinct training frames.",
            "Four reference images predict one canonical Gaussian identity.",
            "Render that identity into four other target poses and calculate losses.",
            "Four groups give 16 target frames per optimizer update; references are resampled.",
            "References and targets are disjoint within each group.",
        ],
        notes="Frozen source: evidence/identity-training-source.py, training loop at line 266. Features are cached per frame. The shared model sees all six identities during training.",
    )
    slide(
        "Our identity setup and its limits",
        bullets=[
            "Frozen Sapiens body and face features; 8,192 canonical Gaussian queries.",
            "Five multimodal blocks, width 1,024; fresh reconstruction network training.",
            "RGB loss balances foreground/background; LPIPS, mask loss and Gaussian priors.",
            "SMPL replaces the paper's MHR template; NeuMan replaces large training datasets.",
            "New frames of six known people do not demonstrate unseen-person generalization.",
        ],
        notes="Exact hyperparameters and deviations: evidence/ours-method.json. No claim of exact LUNA reproduction.",
    )
    slide(
        "Masks and cameras",
        bullets=[
            "NeuMan's provided segmentations define person crops and white reference/GT images.",
            "Test crops also use GT masks: this is oracle-mask preprocessing.",
            "Predictions are never cleaned with the target foreground mask.",
            "LHM++ exports native DPT RGB and its predicted mask; no second compositing.",
            "DPT renders 518 × 518 with unchanged K; retain the original 512 × 512 region.",
        ],
        notes="LHM++ uses the common square inputs, internally resized 512→504 by the released DINO wrapper. This differs from the tall default upstream reference canvas.",
    )
    slide(
        "SMPL to SMPL-X: audited pose conversion",
        bullets=[
            "Both released baselines receive the same converted NeuMan body annotations.",
            "47 fits: 41 test frames plus one reference shape per sequence.",
            "Mean mesh correspondence residual: 3.52 mm.",
            "Mean reprojection difference: 0.71 pixels.",
            "No target RGB fitting; native deformation implementations remain distinct.",
        ],
        notes="evidence/conversion-fits.json; conversion uses official user-supplied mapping assets. Our own model uses SMPL directly.",
    )
    slide(
        "All seven metrics",
        bullets=[
            "PSNR ↑ and L1 ↓: full-crop agreement with the exact target RGB.",
            "Foreground/background L1 ↓: error averaged separately over GT regions.",
            "SSIM ↑: local structure; 11 × 11 Gaussian window, sigma 1.5.",
            "LPIPS-Alex ↓: reference-based perceptual feature distance.",
            "IoU ↑: overlap of predicted and GT masks, threshold 0.5.",
        ],
        notes="Images are lossless 8-bit PNGs in [0,1]. Canonical renders are unscored. Sparse test-frame GIFs do not support claims about MAE/MSJ or temporal smoothness.",
    )
    for title, selected in (
        ("Results: RGB and perceptual agreement", [0, 1, 2, 5]),
        ("Results: background, structure and silhouette", [3, 4, 6]),
    ):
        rows = [["Method / refs"] + [TITLES[i] for i in selected]]
        for method in METHODS:
            scores = method["data"]["mean_over_scenes"]
            rows.append(
                [f"{method['short']} / {method['refs']}"]
                + [f"{scores[KEYS[i]]:.4f}" for i in selected]
            )
        widths = [3.0] + [9.1 / len(selected)] * len(selected)
        slide(
            title,
            table=(rows, widths),
            caption="Equal-weight mean of six scene means. Our model was trained on these identities; released baselines have no local fine-tuning.",
            notes="All unrounded aggregate/per-scene metrics: quantitative/metrics.csv; all 123 method-frame records: quantitative/per-frame.csv.",
        )
    slide(
        "Per-scene RGB and perceptual metrics",
        asset="quantitative/primary-metrics.png",
        caption=delta,
    )
    slide(
        "Per-scene background and silhouette metrics",
        asset="quantitative/additional-metrics.png",
        caption="All six scenes shown. LHM++ IoU uses the DPT mask; LHM and ours use rasterized alpha.",
    )
    for m in media:
        s = m["scene"]
        slide(
            f"{s.capitalize()}: every official test frame",
            asset=m["gif"],
            caption=f"GT | LHM | LHM++ + DPT | our identity. {len(m['targets'])} frames; 300 ms/frame for inspection, not original timing.",
            notes=f"Animated GIF embedded. Separate file: {m['gif']}. Every raw image and predicted mask is included in qualitative/{s}/. Middle test frame is used in the static overview; no best-frame selection.",
        )
    slide(
        "Canonical identity: bike",
        asset="qualitative/bike/canonical.png",
        caption="Same virtual front camera. Canonical ground truth is unavailable; these views do not contribute to metrics.",
    )
    slide(
        "LHM++: canonical identity to posed output",
        asset="qualitative/bike/lhmpp-gt-canonical-posed.gif",
        caption="Ground truth | canonical identity | native SMPL-X posing + DPT. One reconstruction reused across targets.",
    )
    slide(
        "Our identity: canonical identity to LBS output",
        asset="qualitative/bike/ours-gt-canonical-lbs.gif",
        caption="Ground truth | canonical identity | SMPL/LBS render. This is the teacher posing path, not neural animation.",
    )
    slide(
        "Why visual quality and metrics can disagree",
        bullets=[
            "LHM++ looks smoother than ours, but head direction and body extent can differ from GT.",
            "Versus LHM: foreground L1 worsens (0.1301 → 0.1620); background improves (0.0099 → 0.0030).",
            "Full-crop L1 mixes both regions. PSNR weights large errors more through squared error.",
            "Our lower errors coexist with rough faces, clothing edges and visible Gaussian splats.",
            "LPIPS also compares to target images; these scores alone do not rank realism.",
        ],
        notes="Fixed middle-frame observations: LHM++ gives coherent parkinglot hair/headphones and jogging jacket folds; bike body extent and lab/seattle head directions still mismatch GT. Every test frame is supplied. Causal attribution requires ablations.",
    )
    slide(
        "What this comparison establishes",
        bullets=[
            "Measured all three methods on the same 41 official NeuMan test frames.",
            delta,
            "LHM++ and ours have matching reference counts; their training exposure differs.",
            "Ours leads six aggregate metrics; LHM++ has the lowest background L1.",
            "These results do not establish unseen-person performance or full LUNA animation quality.",
        ],
    )
    slide(
        "Next experiments",
        bullets=[
            "Evaluate new identities excluded from our identity training.",
            "Separate geometry/alignment error from appearance detail with explicit diagnostics.",
            "Ablate reference count and training data under a declared validation protocol.",
            "Train and evaluate the neural animator separately from LBS identity reconstruction.",
            "Use continuous trajectories and timestamps for temporal metrics.",
        ],
    )
    slide(
        "Reproducibility and deliverables",
        bullets=[
            "NeuMan-final-report.pptx: editable text, tables, embedded comparison GIFs.",
            "REPORT.html: complete report and local gallery.",
            "qualitative/: all scenes, 18 GIFs, raw predictions, masks and canonical views.",
            "quantitative/: seven metrics, scene means, every frame and editable SVG plots.",
            "evidence/: pinned models/source, geometry checks, configuration and adapters.",
        ],
        notes="Sources: https://arxiv.org/html/2606.31981v2 ; https://arxiv.org/html/2503.10625v1 ; https://arxiv.org/html/2506.13766v2 . LHM++ implementation is called without copying learned modules; source, weights and body assets have separate licenses.",
    )
    prs.save(out / "NeuMan-final-report.pptx")
    (out / "SLIDES.md").write_text("# Final presentation outline\n\n" + "\n".join(outline))
    (out / "slide-specs.json").write_text(json.dumps(specs, indent=2))


if __name__ == "__main__":
    main()
