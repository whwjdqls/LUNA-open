"""Check the delivered report, metrics, GIFs and PPTX; create portable previews/ZIP.

Slide previews are rendered from PPTX shapes with Pillow, not Microsoft Office.
They check content/layout and embedded media; Office playback remains untested.
"""

import argparse
import csv
import hashlib
import io
import json
import os
import socket
import zipfile
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlsplit

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

KEYS = ["psnr", "l1", "foreground_l1", "background_l1", "ssim", "lpips", "mask_iou"]


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(8 * 1024**2), b""):
            h.update(chunk)
    return h.hexdigest()


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attrs):
        for key, value in attrs:
            if key in ("href", "src"):
                self.links.append(value)


def slide_previews(prs, out):
    dpi = 120
    scale = dpi / 914400
    width, height = round(prs.slide_width * scale), round(prs.slide_height * scale)
    previews, overflow = [], []

    def draw_text(draw, frame, x, y, w, h, slide_no, shape_name):
        y += round(frame.margin_top * scale)
        bottom = y + h - round((frame.margin_top + frame.margin_bottom) * scale)
        x += round(frame.margin_left * scale)
        usable = w - round((frame.margin_left + frame.margin_right) * scale)
        for pi, p in enumerate(frame.paragraphs):
            if not p.text:
                continue
            size = p.font.size.pt if p.font.size else 18
            bold = bool(p.font.bold)
            path = (
                "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans" + ("-Bold" if bold else "") + ".ttf"
            )
            font = ImageFont.truetype(path, max(1, round(size * dpi / 72)))
            color = "#243448"
            if p.font.color.type is not None:
                try:
                    color = "#" + str(p.font.color.rgb)
                except AttributeError:
                    pass
            words, lines, line = p.text.split(), [], ""
            for word in words:
                candidate = (line + " " + word).strip()
                if line and draw.textlength(candidate, font=font) > usable:
                    lines.append(line)
                    line = word
                else:
                    line = candidate
            if line:
                lines.append(line)
            line_height = round(size * dpi / 72 * 1.18)
            for line in lines:
                if y + line_height > bottom + 6:
                    overflow.append(dict(slide=slide_no, shape=shape_name, text=line))
                draw.text((x, y), line, fill=color, font=font, anchor="lt")
                y += line_height
            if pi < len(frame.paragraphs) - 1:
                y += round((p.space_after.pt if p.space_after else 0) * dpi / 72)

    for number, slide in enumerate(prs.slides, 1):
        canvas = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(canvas)
        for shape in slide.shapes:
            x, y, w, h = [
                round(v * scale) for v in (shape.left, shape.top, shape.width, shape.height)
            ]
            if hasattr(shape, "image"):
                with Image.open(io.BytesIO(shape.image.blob)) as picture:
                    canvas.paste(
                        picture.convert("RGB").resize((w, h), Image.Resampling.LANCZOS), (x, y)
                    )
            elif shape.has_table:
                ty = y
                for row in shape.table.rows:
                    tx = x
                    rh = round(row.height * scale)
                    for col, cell in zip(shape.table.columns, row.cells):
                        cw = round(col.width * scale)
                        fill = "#" + str(cell.fill.fore_color.rgb)
                        draw.rectangle((tx, ty, tx + cw, ty + rh), fill=fill, outline="#ccd4dc")
                        draw_text(draw, cell.text_frame, tx, ty, cw, rh, number, "table cell")
                        tx += cw
                    ty += rh
            elif shape.has_text_frame:
                draw_text(draw, shape.text_frame, x, y, w, h, number, shape.name)
        path = out / "slide-previews" / f"slide-{number:02d}.png"
        canvas.save(path)
        previews.append(canvas)
    if overflow:
        (out / "slide-previews/layout-overflow.json").write_text(json.dumps(overflow, indent=2))
        raise AssertionError(f"Slide text may overflow: {overflow[:5]}")
    previews[0].save(
        out / "NeuMan-final-report-preview.pdf",
        save_all=True,
        append_images=previews[1:],
        resolution=dpi,
    )
    for start in range(0, len(previews), 9):
        contact = Image.new("RGB", (1200, 3 * 245), "#dce4ec")
        for i, picture in enumerate(previews[start : start + 9]):
            picture = picture.resize((390, 219), Image.Resampling.LANCZOS)
            contact.paste(picture, ((i % 3) * 400 + 5, (i // 3) * 245 + 20))
        contact.save(out / "slide-previews" / f"contact-{start // 9 + 1}.png")


def main():
    if not os.environ.get("SLURM_JOB_ID") or "login" in socket.gethostname():
        raise RuntimeError("Slurm compute node required")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    out = args.output
    protocol = json.loads((out / "evidence/protocol.json").read_text())
    expected = {(s, f) for s, i in protocol["scenes"].items() for f in i["targets"]}
    assert len(expected) == 41
    links_checked = 0
    for page in out.rglob("*.html"):
        parser = Links()
        parser.feed(page.read_text())
        for value in parser.links:
            parsed = urlsplit(value)
            if parsed.scheme or parsed.netloc:
                continue
            assert (page.parent / unquote(parsed.path)).exists(), (page, value)
            links_checked += 1
    metric_checks = {}
    csv_rows = list(csv.DictReader((out / "quantitative/metrics.csv").open()))
    assert len(csv_rows) == 21
    assert len(list(csv.DictReader((out / "quantitative/per-frame.csv").open()))) == 123
    for method in ("lhm", "lhmpp", "ours"):
        data = json.loads((out / f"evidence/{method}-metrics.json").read_text())
        assert len(data["frames"]) == 41
        assert {(r["scene"], r["frame"]) for r in data["frames"]} == expected
        assert data["method"]["manifest_sha256"] == protocol["manifest_sha256"]
        for scene in protocol["scenes"]:
            rows = [r["metrics"] for r in data["frames"] if r["scene"] == scene]
            for key in KEYS:
                assert abs(np.mean([r[key] for r in rows]) - data["per_scene"][scene][key]) < 1e-10
        for key in KEYS:
            assert (
                abs(
                    np.mean([r[key] for r in data["per_scene"].values()])
                    - data["mean_over_scenes"][key]
                )
                < 1e-10
            )
        max_errors = dict(l1=0.0, psnr=0.0)
        for row in data["frames"]:
            folder = out / "qualitative" / row["scene"]
            with Image.open(folder / method / "rgb" / row["frame"]) as im:
                assert im.size == (512, 512)
                pred = np.asarray(im.convert("RGB"), dtype=np.float64) / 255
            with Image.open(folder / "ground-truth" / row["frame"]) as im:
                gt = np.asarray(im.convert("RGB"), dtype=np.float64) / 255
            with Image.open(folder / method / "alpha" / row["frame"]) as im:
                assert im.size == (512, 512)
            # Protocol GT PNG is quantized; the evaluator reads float GT from
            # the dataset. Bound this independent check by 1/255 RGB rounding.
            l1 = np.abs(pred - gt).mean()
            mse = ((pred - gt) ** 2).mean()
            psnr = -10 * np.log10(max(mse, 1e-10))
            max_errors["l1"] = max(max_errors["l1"], abs(l1 - row["metrics"]["l1"]))
            max_errors["psnr"] = max(max_errors["psnr"], abs(psnr - row["metrics"]["psnr"]))
        assert max_errors["l1"] < 1 / 255 + 1e-6
        assert max_errors["psnr"] < 0.15
        metric_checks[method] = dict(
            all_seven_aggregates_match=True, independent_png_check=max_errors
        )
    for row in csv_rows:
        method = (
            "lhmpp"
            if row["method"].startswith("LHM++")
            else "lhm"
            if row["method"].startswith("LHM-")
            else "ours"
        )
        data = json.loads((out / f"evidence/{method}-metrics.json").read_text())
        correct = (
            data["mean_over_scenes"] if row["scene"] == "macro" else data["per_scene"][row["scene"]]
        )
        for key in KEYS:
            assert abs(float(row[key]) - correct[key]) < 1e-12
    gifs, image_count = [], 0
    for path in (out / "qualitative").rglob("*"):
        if path.suffix not in (".png", ".gif"):
            continue
        with Image.open(path) as im:
            if path.suffix == ".gif":
                count = len(protocol["scenes"][path.parent.name]["targets"])
                assert im.n_frames == count, (path, im.n_frames, count)
                for i in range(im.n_frames):
                    im.seek(i)
                    im.load()
                gifs.append(dict(path=str(path.relative_to(out)), frames=count))
            else:
                im.verify()
        image_count += 1
    assert len(gifs) == 18
    deck = out / "NeuMan-final-report.pptx"
    prs = Presentation(deck)
    assert len(prs.slides) == 27, len(prs.slides)
    for slide in prs.slides:
        assert slide.notes_slide.notes_text_frame.text.strip()
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 100
            assert shape.top + shape.height <= prs.slide_height + 100
    embedded_gifs = []
    with zipfile.ZipFile(deck) as archive:
        assert archive.testzip() is None
        blobs = {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        }
        for scene in protocol["scenes"]:
            assert digest(out / f"qualitative/{scene}/comparison.gif") in blobs
        for name in archive.namelist():
            if name.startswith("ppt/media/") and name.endswith(".gif"):
                with Image.open(io.BytesIO(archive.read(name))) as im:
                    assert im.n_frames > 1
                    embedded_gifs.append(dict(file=name, frames=im.n_frames))
    assert len(embedded_gifs) == 8
    slide_previews(prs, out)
    receipt = dict(
        checked_at=datetime.now().astimezone().isoformat(),
        host=socket.gethostname(),
        job=os.environ["SLURM_JOB_ID"],
        scenes=6,
        test_frames_per_method=41,
        methods=3,
        metrics=KEYS,
        metric_checks=metric_checks,
        local_html_links_checked=links_checked,
        images_decoded=image_count,
        gifs=gifs,
        slides=len(prs.slides),
        embedded_gifs=embedded_gifs,
        all_gif_frames_decoded=True,
        slide_bounds_and_text_fit_checked=True,
        pptx_crc_passed=True,
        preview_scope="Pillow rendering of actual PPTX shapes; Microsoft PowerPoint/LibreOffice playback not executed",
    )
    (out / "PACKAGE_CHECKS.json").write_text(json.dumps(receipt, indent=2) + "\n")
    paths = sorted(p for p in out.rglob("*") if p.is_file() and p.name != "FILE_MANIFEST.csv")
    with (out / "FILE_MANIFEST.csv").open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["path", "bytes", "sha256"])
        writer.writeheader()
        for path in paths:
            writer.writerow(
                dict(
                    path=str(path.relative_to(out)), bytes=path.stat().st_size, sha256=digest(path)
                )
            )
    package = out.with_suffix(".zip")
    with zipfile.ZipFile(
        package, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=4
    ) as archive:
        for path in sorted(out.rglob("*")):
            if path.is_file():
                archive.write(path, str(Path(out.name) / path.relative_to(out)))
    with zipfile.ZipFile(package) as archive:
        assert archive.testzip() is None
    package_hash = digest(package)
    package.with_suffix(".zip.sha256").write_text(f"{package_hash}  {package.name}\n")
    audit = dict(
        package=str(package),
        bytes=package.stat().st_size,
        sha256=package_hash,
        deck=str(deck),
        slides=len(prs.slides),
        archive_crc_passed=True,
    )
    (out.parent / (out.name + "-audit.json")).write_text(json.dumps(audit, indent=2) + "\n")
    print(json.dumps(audit, indent=2), flush=True)


if __name__ == "__main__":
    main()
