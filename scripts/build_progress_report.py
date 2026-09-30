"""Build a portable scientific report, plots, gallery and editable starter deck."""

import csv
import html
import json
import os
import re
import shutil
import socket
from datetime import datetime
from pathlib import Path

import markdown2
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

WORK = Path("/scratch2/whwjdqls99/LUNA-open")
REPO = Path("/home/whwjdqls99/LUNA-open")
OUT = WORK / "reports/luna-progress-20260928"
CHARTS = OUT / "quantitative/charts"
BLUE = "#2166ac"
GREEN = "#16846b"
RED = "#bf4b45"


def read(path):
    return json.loads(Path(path).read_text())


def csv_write(path, rows):
    with Path(path).open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def md_table(headers, rows):
    return "\n".join(
        ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
        + ["| " + " | ".join(str(v) for v in row) + " |" for row in rows]
    )


def save_chart(fig, name):
    fig.savefig(CHARTS / f"{name}.png", dpi=210, bbox_inches="tight", facecolor="white")
    fig.savefig(CHARTS / f"{name}.svg", bbox_inches="tight", facecolor="white")
    plt.close(fig)


def font(size):
    return ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", size)


def montage(paths, labels, destination, title, footer):
    size = 512
    canvas = Image.new("RGB", (len(paths) * size + 32, 620), "#eef1f5")
    draw = ImageDraw.Draw(canvas)
    draw.text((16, 10), title, font=font(23), fill="#142940")
    for i, (path, label) in enumerate(zip(paths, labels, strict=True)):
        draw.text((16 + i * size, 46), label, font=font(18), fill="#142940")
        with Image.open(path) as picture:
            canvas.paste(
                picture.convert("RGB").resize((size, size), Image.Resampling.LANCZOS),
                (16 + i * size, 76),
            )
    draw.text((16, 595), footer, font=font(15), fill="#465465")
    canvas.save(destination)


def main():
    if not os.environ.get("SLURM_JOB_ID") or "login" in socket.gethostname():
        raise RuntimeError("Use a CPU compute allocation")
    snapshot = read(OUT / "evidence/report-snapshot.json")
    bench = read(OUT / "evidence/audits/batching-report.json")
    # Complete the portable evidence collection without copying model weights.
    for label, source in [
        ("11_geometry_audit", WORK / "outputs/smpl-audit-2336974"),
        ("12_identity_pilot", WORK / "outputs/identity-pilot-2336972"),
    ]:
        dest = OUT / "qualitative" / label
        dest.mkdir(exist_ok=True)
        for path in source.iterdir():
            if path.is_file() and path.suffix in {".png", ".jpg", ".json"}:
                shutil.copy2(path, dest / path.name)
    extras = [
        WORK / "assets/download-verification.json",
        WORK / "acquisition-summary.json",
        WORK / "runs/neuman/provenance/environment.json",
        WORK / "outputs/identity-resume-diagnostic-2336972.json",
        WORK / "outputs/gpu-2336972/animator-articulation-6000-summary.json",
    ]
    for source in extras:
        shutil.copy2(source, OUT / "evidence/audits" / source.name)
    for name in (
        "model.py",
        "smpl.py",
        "training.py",
        "identity_training.py",
        "identity_batching.py",
        "rendering.py",
        "losses.py",
        "metrics.py",
        "features.py",
        "pipeline.py",
        "provenance.py",
    ):
        shutil.copy2(REPO / "src/luna_open" / name, OUT / "evidence/source" / name)
    shutil.copy2(REPO / "src/luna_open/data/neuman.py", OUT / "evidence/source/neuman.py")
    for name in (
        "prepare_progress_report.py",
        "render_progress_comparison.py",
        "build_progress_report.py",
    ):
        shutil.copy2(REPO / "scripts" / name, OUT / "evidence/source" / name)
    slidefig = OUT / "qualitative/13_slide_figures"
    slidefig.mkdir(exist_ok=True)
    pilot = OUT / "qualitative/12_identity_pilot"
    montage(
        [pilot / "target.png", pilot / "initial-rgb.png", pilot / "final-rgb.png"],
        ["Ground truth", "Initial prediction", "After 40 updates"],
        slidefig / "fixed-frame-pilot.png",
        "Fixed-frame identity learning pilot",
        "Training-target diagnostic; strict continuation assertion failed separately",
    )
    diag = OUT / "qualitative/05_identity_diagnosis/bike"
    montage(
        [diag / "target-0.png", diag / "baseline-0.png", diag / "direct-fit-0.png"],
        ["Ground truth", "Original encoder output", "Direct Gaussian fit, 128 steps"],
        slidefig / "direct-fit-diagnostic.png",
        "Representation-capacity diagnostic: bike training target",
        "Direct fitting on this image; not an encoder or held-out generalization result",
    )
    anim = OUT / "qualitative/03_neural_animator_5000_failure/citron"
    montage(
        [anim / "target-00007.png", anim / "identity-00007.png", anim / "animator-00007.png"],
        ["Ground truth / driver", "Identity + fitted SMPL/LBS", "Neural animator, step 5,000"],
        slidefig / "animator-failure.png",
        "Original neural animator: selected checkpoint",
        "Held-out test frame; LBS articulation and neural animation are different outputs",
    )

    originals = {}
    model_rows = []
    for stage in ("identity", "animator"):
        for split in ("val", "test"):
            result = read(OUT / f"quantitative/raw/original_{stage}/{split}-metrics.json")
            originals[(stage, split)] = result
            model_rows.append(
                dict(
                    model=f"original_{stage}",
                    split=split,
                    checkpoint_update=10000 if stage == "identity" else 5000,
                    frames=len(result["frames"]),
                    **result["mean_over_scenes"],
                )
            )
    new_step = snapshot["models"]["retrained_identity"]["update"]
    new_result = read(OUT / f"quantitative/raw/retrained_identity/val-{new_step:06d}.json")
    model_rows.append(
        dict(
            model="retrained_identity",
            split="val",
            checkpoint_update=new_step,
            frames=len(new_result["frames"]),
            **{k: v for k, v in new_result["mean_over_scenes"].items() if k != "foreground_l1"},
        )
    )
    csv_write(OUT / "quantitative/model-summary.csv", model_rows)
    same = {
        label: read(OUT / f"quantitative/{label}-same-frame-val.json")
        for label in ("original_identity", "retrained_identity")
    }
    expected = {
        "original_identity": originals[("identity", "val")],
        "retrained_identity": new_result,
    }
    for label in same:
        assert len(same[label]["frames"]) == 44
        assert {(r["scene"], r["frame"]) for r in same[label]["frames"]} == {
            (r["scene"], r["frame"]) for r in expected[label]["frames"]
        }
        for key in ("psnr", "l1", "ssim", "mask_iou", "lpips"):
            assert (
                abs(same[label]["mean_over_scenes"][key] - expected[label]["mean_over_scenes"][key])
                < 1e-7
            ), (label, key)
    same_rows = [
        dict(model=label, update=same[label]["update"], **same[label]["mean_over_scenes"])
        for label in same
    ]
    csv_write(OUT / "quantitative/same-frame-summary.csv", same_rows)
    per_scene = [
        dict(model=label, scene=scene, **metrics)
        for label, result in same.items()
        for scene, metrics in result["per_scene"].items()
    ]
    csv_write(OUT / "quantitative/per-scene-validation.csv", per_scene)
    curves = {}
    for label in ("original_identity", "original_animator", "retrained_identity"):
        records = [
            json.loads(line)
            for line in (OUT / f"quantitative/raw/{label}/train.jsonl").read_text().splitlines()
            if line.strip()
        ]
        assert all(row["update"] == i + 1 for i, row in enumerate(records))
        curves[label] = [
            dict(model=label, update=r["update"], lpips=r["validation_lpips"])
            for r in records
            if "validation_lpips" in r
        ]
    csv_write(
        OUT / "quantitative/validation-curves.csv",
        [row for curve in curves.values() for row in curve],
    )
    batches = [dict(variant=mode, **values) for mode, values in bench["summary"].items()]
    csv_write(OUT / "quantitative/batching-summary.csv", batches)
    splits = [dict(scene=scene, **values) for scene, values in snapshot["splits"].items()]
    csv_write(OUT / "quantitative/dataset-split.csv", splits)
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 12,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.titleweight": "bold",
        }
    )
    # Dataset chart.
    fig, ax = plt.subplots(figsize=(10, 4.4))
    left = np.zeros(len(splits))
    for key, color in [("train", BLUE), ("val", GREEN), ("test", "#e6a447")]:
        values = np.array([r[key] for r in splits])
        ax.barh([r["scene"] for r in splits], values, left=left, label=key, color=color)
        left += values
    ax.set(xlabel="Frames", title="Official NeuMan split: 344 train / 44 validation / 41 test")
    ax.invert_yaxis()
    ax.legend(ncol=3, loc="lower right")
    save_chart(fig, "dataset-split")
    # Identity curves: exact rows in frozen logs.
    fig, ax = plt.subplots(figsize=(10, 4.8))
    for label, color, title in [
        ("original_identity", BLUE, "Original identity"),
        ("retrained_identity", GREEN, "Revised identity (ongoing)"),
    ]:
        data = curves[label]
        ax.plot(
            [r["update"] for r in data],
            [r["lpips"] for r in data],
            label=title,
            color=color,
            marker=".",
            linewidth=2,
        )
    ax.scatter([new_step], [new_result["mean_over_scenes"]["lpips"]], color=GREEN, s=85, zorder=4)
    ax.set(
        xlabel="Optimizer updates (effective batch 16)",
        ylabel="Validation LPIPS ↓",
        title="Identity reconstruction: held-out frames of six seen subjects",
    )
    ax.grid(alpha=0.2)
    ax.legend()
    save_chart(fig, "identity-validation")
    # Per-scene values from new, identical-frame evaluation.
    fig, ax = plt.subplots(figsize=(10, 4.8))
    scenes = list(same["original_identity"]["per_scene"])
    x = np.arange(len(scenes))
    for offset, label, color, title in [
        (-0.19, "original_identity", BLUE, "Original 10k"),
        (0.19, "retrained_identity", GREEN, f"Revised {new_step:,}"),
    ]:
        ax.bar(
            x + offset,
            [same[label]["per_scene"][s]["lpips"] for s in scenes],
            width=0.36,
            label=title,
            color=color,
        )
    ax.set(
        xticks=x,
        xticklabels=scenes,
        ylabel="Validation LPIPS ↓",
        title="Same references, target frames, crops and SMPL/LBS poses",
    )
    ax.legend()
    save_chart(fig, "per-scene-lpips")
    # Animator failure curve.
    fig, ax = plt.subplots(figsize=(10, 4.6))
    data = curves["original_animator"]
    ax.plot(
        [r["update"] for r in data], [r["lpips"] for r in data], color=RED, marker="o", markersize=4
    )
    ax.axvspan(0, 1000, color="#e8edf2", label="Global-motion warmup")
    ax.scatter(
        [5000],
        [originals[("animator", "val")]["mean_over_scenes"]["lpips"]],
        s=90,
        color="#172940",
        label="Selected checkpoint (5k)",
    )
    ax.set(
        xlabel="Animator updates",
        ylabel="Validation LPIPS ↓",
        title="Original animator: completion did not repair local articulation",
    )
    ax.legend()
    ax.grid(alpha=0.2)
    save_chart(fig, "animator-validation")
    # Target FG metric over time.
    newvals = [
        read(p) for p in sorted((OUT / "quantitative/raw/retrained_identity").glob("val-*.json"))
    ]
    fg_rows = [
        dict(update=v["update"], foreground_l1=v["mean_over_scenes"]["foreground_l1"])
        for v in newvals
    ]
    csv_write(OUT / "quantitative/retraining-foreground-l1.csv", fg_rows)
    fig, ax = plt.subplots(figsize=(10, 4.4))
    ax.plot([r["update"] for r in fg_rows], [r["foreground_l1"] for r in fg_rows], color=GREEN)
    ax.axhline(
        same["original_identity"]["mean_over_scenes"]["foreground_l1"],
        color=BLUE,
        linestyle="--",
        label="Original 10k (same-frame re-evaluation)",
    )
    ax.set(
        xlabel="Revised identity updates",
        ylabel="Foreground RGB L1 ↓",
        title="Foreground reconstruction error",
    )
    ax.legend()
    ax.grid(alpha=0.2)
    save_chart(fig, "foreground-error")
    # Batch benchmark.
    fig, axes = plt.subplots(1, 2, figsize=(12, 5.2))
    labels = [
        "Current serial",
        "Deferred metrics",
        "1 ID + batched losses",
        "2 IDs + serial losses",
        "2 IDs + batched losses",
    ]
    y = np.arange(5)
    for ax, key, title in [
        (axes[0], "median_seconds", "Seconds per update ↓"),
        (axes[1], "peak_allocated_gib", "Peak allocated GPU memory (GiB)"),
    ]:
        values = [r[key] for r in batches]
        ax.barh(y, values, color=[BLUE, BLUE, GREEN, "#8c7ca5", GREEN])
        ax.set_yticks(y, labels if ax is axes[0] else [""] * 5)
        ax.invert_yaxis()
        ax.set_title(title)
        for i, v in enumerate(values):
            ax.text(
                v + 0.02,
                i,
                f"{v:.3f}" if key == "median_seconds" else f"{v:.2f}",
                va="center",
                fontsize=10,
            )
        ax.set_xlim(0, max(values) * 1.22)
    fig.suptitle(
        "Separate RTX 4090: effective batch 16; 20 measured updates per variant", fontsize=14
    )
    fig.tight_layout()
    save_chart(fig, "batching")
    # Scientific pipeline diagram, made from our code interfaces.
    fig, ax = plt.subplots(figsize=(13, 5.5))
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 5.5)
    ax.axis("off")

    def box(x, y, w, h, text, color):
        ax.add_patch(
            FancyBboxPatch(
                (x, y), w, h, boxstyle="round,pad=.07", facecolor=color, edgecolor="#96a5b5"
            )
        )
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=11)

    def arrow(a, b, color="#5a6d80"):
        ax.add_patch(
            FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=18, linewidth=1.8, color=color)
        )

    box(0.2, 3.45, 2.0, 1.2, "4 reference images\nFrozen body + face\nfeatures", "#e9f0fa")
    box(
        2.8,
        3.45,
        2.3,
        1.2,
        "Identity transformer\n8,192 queries\n5 blocks × width 1,024",
        "#e9f0fa",
    )
    box(5.7, 3.45, 2.1, 1.2, "Canonical Gaussians\n+ identity tokens", "#e4f3eb")
    box(
        9.05,
        3.45,
        3.3,
        1.2,
        "SMPL/LBS + target camera\nRendered RGB / mask\nIdentity supervision",
        "#e4f3eb",
    )
    box(5.7, 0.8, 2.1, 1.25, "Driving image\nFrozen DINOv3", "#fff1db")
    box(
        9.05,
        0.8,
        3.3,
        1.25,
        "Neural animator\nGlobal motion + local deltas\nImage-driven posed Gaussians",
        "#ffe8e5",
    )
    for a, b in [
        ((2.2, 4.05), (2.8, 4.05)),
        ((5.1, 4.05), (5.7, 4.05)),
        ((7.8, 4.05), (9.05, 4.05)),
        ((7.8, 1.4), (9.05, 1.4)),
        ((6.75, 3.45), (9.2, 2.05)),
    ]:
        arrow(a, b)
    ax.text(
        6.5,
        5.13,
        "Identity reconstruction and neural animation have different supervision paths",
        ha="center",
        fontsize=15,
        weight="bold",
    )
    ax.text(
        0.25,
        0.22,
        "Green path: identity renders use supplied fitted poses. Red path: learned animator currently fails local articulation.",
        fontsize=11,
    )
    save_chart(fig, "pipeline")

    def model_table(stage):
        return md_table(
            [
                "Split",
                "Selected update",
                "Frames",
                "PSNR ↑",
                "L1 ↓",
                "SSIM ↑",
                "Mask IoU ↑",
                "LPIPS ↓",
            ],
            [
                [
                    r["split"],
                    r["checkpoint_update"],
                    r["frames"],
                    f"{r['psnr']:.4f}",
                    f"{r['l1']:.6f}",
                    f"{r['ssim']:.6f}",
                    f"{r['mask_iou']:.6f}",
                    f"{r['lpips']:.6f}",
                ]
                for r in model_rows
                if r["model"] == f"original_{stage}"
            ],
        )

    old_lpips = originals[("identity", "val")]["mean_over_scenes"]["lpips"]
    new_lpips = new_result["mean_over_scenes"]["lpips"]
    progress_steps = {250, 1000, 3000, 5000, 7500, 10000, new_step, newvals[-1]["update"]}
    progress_table = md_table(
        ["Revised update", "Validation LPIPS ↓", "Foreground L1 ↓"],
        [
            [
                v["update"],
                f"{v['mean_over_scenes']['lpips']:.6f}",
                f"{v['mean_over_scenes']['foreground_l1']:.6f}",
            ]
            for v in newvals
            if v["update"] in progress_steps
        ],
    )
    same_table = md_table(
        ["Model", "Update", "PSNR ↑", "SSIM ↑", "Mask IoU ↑", "LPIPS ↓", "Foreground L1 ↓"],
        [
            [
                r["model"],
                r["update"],
                f"{r['psnr']:.4f}",
                f"{r['ssim']:.6f}",
                f"{r['mask_iou']:.6f}",
                f"{r['lpips']:.6f}",
                f"{r['foreground_l1']:.6f}",
            ]
            for r in same_rows
        ],
    )
    batch_table = md_table(
        ["Variant", "Seconds/update", "Targets/second", "Peak GiB", "Speedup"],
        [
            [
                r["variant"],
                f"{r['median_seconds']:.3f}",
                f"{r['targets_per_second']:.3f}",
                f"{r['peak_allocated_gib']:.2f}",
                f"{r['speedup_over_serial']:.3f}×",
            ]
            for r in batches
        ],
    )
    replacements = {
        "PROGRESS": f"{snapshot['training_update_at_snapshot']:,}",
        "NEW_STEP": f"{new_step:,}",
        "NEW_LPIPS": f"{new_lpips:.6f}",
        "IMPROVEMENT": f"{100 * (old_lpips - new_lpips) / old_lpips:.1f}",
        "SNAPSHOT_TIME": snapshot["created_at"],
        "SPLIT_TABLE": md_table(
            ["Sequence", "Train", "Validation", "Test"],
            [[r["scene"], r["train"], r["val"], r["test"]] for r in splits]
            + [["Total", 344, 44, 41]],
        ),
        "ORIGINAL_TABLE": model_table("identity"),
        "ANIMATOR_TABLE": model_table("animator"),
        "IDENTITY_PROGRESS_TABLE": progress_table,
        "SAME_FRAME_TABLE": same_table,
        "BATCH_TABLE": batch_table,
    }
    report = (REPO / "scripts/report_assets/progress-report.md").read_text()
    for key, value in replacements.items():
        report = report.replace("{{" + key + "}}", value)
    assert not re.search(r"\{\{[A-Z_]+\}\}", report)
    (OUT / "REPORT.md").write_text(report)
    css = "body{font:16px/1.65 system-ui,sans-serif;color:#193048;max-width:1160px;margin:40px auto;padding:0 24px}h1,h2,h3{line-height:1.25}h2{margin-top:2.5em;border-top:1px solid #ccd5df;padding-top:1em}a{color:#1765a4}img{max-width:100%;height:auto}table{border-collapse:collapse;width:100%;font-size:14px}td,th{border:1px solid #d2dce5;padding:9px;text-align:left}th{background:#edf2f7}pre{overflow:auto;background:#edf2f7;padding:18px}blockquote{border-left:4px solid #16846b;padding:4px 20px;background:#f1f8f4}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:20px}figure{margin:15px 0}figcaption{font-size:14px}details{padding:14px;background:#f3f6f9;margin:16px 0}li{margin:5px 0}"
    rendered = markdown2.markdown(
        report, extras=["tables", "fenced-code-blocks", "header-ids", "toc"]
    )
    (OUT / "REPORT.html").write_text(
        f'<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>LUNA-open progress report</title><style>{css}</style><nav><a href="gallery.html">Qualitative gallery</a> · <a href="LUNA-open-progress.pptx">PowerPoint</a> · <a href="REPORT.md">Markdown report</a></nav>{rendered}</html>'
    )

    sections = [
        (
            "same-frame-comparison",
            "Original versus revised identity on identical validation frames",
            ["10_same_frame_identity_comparison"],
            "Both columns use fitted SMPL/LBS. Revised checkpoint " + str(new_step) + ".",
        ),
        (
            "original-identity",
            "Original identity: canonical views, all train/test frames and GIFs",
            ["01_original_identity_all_frames", "02_original_identity_gt_canonical_lbs_gifs"],
            "Completed original identity at 10,000 updates. Posed outputs use annotated SMPL/LBS.",
        ),
        (
            "neural-animation",
            "Original neural animation: articulation failure",
            ["03_neural_animator_5000_failure", "04_canonical_and_neural_animation_5000"],
            "Selected original animator at 5,000 updates. These examples show a failure, not successful animation.",
        ),
        (
            "diagnostics",
            "Diagnostics and isolated fitting probes",
            [
                "05_identity_diagnosis",
                "06_animator_normalization_probe",
                "07_animator_fresh_local_probe",
                "12_identity_pilot",
                "13_slide_figures",
            ],
            "Training-target fitting and debugging experiments; do not treat them as held-out model improvements.",
        ),
        (
            "verification",
            "Geometry, smoke and retraining progression",
            ["08_smoke_example", "09_retraining_validation_progress", "11_geometry_audit"],
            "Smoke outputs include synthetic inputs. Validation previews use identity/LBS rendering.",
        ),
    ]
    body = [
        '<h1>LUNA-open qualitative results</h1><p><a href="REPORT.html">Full report</a> · <a href="LUNA-open-progress.pptx">Starter PowerPoint</a></p><p>All media is local to this folder. Open an image for full resolution. GIF playback is illustrative; canonical and fitted-pose outputs are explicitly distinguished from neural animation.</p>'
    ]
    media_inventory = []
    for anchor, title, folders, caption in sections:
        body.append(
            f'<section id="{anchor}"><h2>{html.escape(title)}</h2><p>{html.escape(caption)}</p><div class="grid">'
        )
        files = [
            p
            for folder in folders
            for p in sorted((OUT / "qualitative" / folder).rglob("*"))
            if p.is_file() and p.suffix.lower() in {".png", ".jpg", ".gif", ".mp4"}
        ]
        featured = [
            p
            for p in files
            if (
                p.parent == OUT / "qualitative/10_same_frame_identity_comparison"
                and p.name.endswith(("-comparison.png", "-gt-canonical-lbs.gif"))
            )
            or p.name
            in {
                "train-test-1.jpg",
                "train-test-2.jpg",
                "comparison.jpg",
                "optimized-overlays.png",
                "fixed-frame-pilot.png",
                "direct-fit-diagnostic.png",
                "animator-failure.png",
            }
            or (p.parent.name in {"test"} and p.suffix == ".gif")
            or (p.name == "held-out-sequence.gif")
        ]
        if not featured:
            featured = files[:6]
        for p in featured:
            rel = p.relative_to(OUT).as_posix()
            body.append(
                f'<figure><a href="{rel}"><img loading="lazy" src="{rel}" alt="{html.escape(p.name)}"></a><figcaption>{html.escape(str(p.relative_to(OUT / "qualitative")))}</figcaption></figure>'
            )
        body.append(
            f"</div><details><summary>All {len(files)} media files in this section</summary><ul>"
        )
        for p in files:
            rel = p.relative_to(OUT).as_posix()
            body.append(
                f'<li><a href="{rel}">{html.escape(str(p.relative_to(OUT / "qualitative")))}</a></li>'
            )
            media_inventory.append(dict(section=anchor, path=rel, description=caption))
        body.append("</ul></details></section>")
    (OUT / "gallery.html").write_text(
        f'<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>LUNA-open qualitative gallery</title><style>{css}</style>'
        + "".join(body)
        + "</html>"
    )
    csv_write(OUT / "qualitative/MEDIA_INDEX.csv", media_inventory)

    # Editable slide text; figures remain faithful static research images.
    q = "qualitative/10_same_frame_identity_comparison/"
    slides = [
        dict(
            title="LUNA-open: progress and open problems",
            bullets=[
                "Independent NeuMan / SMPL development reimplementation",
                "Yonsei RTX 4090 experiments • September 28, 2026",
                f"Retraining snapshot: {snapshot['training_update_at_snapshot']:,}/20,000 updates; selected checkpoint {new_step:,}",
                "Canonical reconstruction improved; neural animation remains unresolved.",
            ],
        ),
        dict(
            title="What we have completed",
            bullets=[
                "Prepared and verified local data, assets, geometry and frozen features.",
                "Completed 10k identity + 10k animator development training.",
                "Diagnosed animator collapse and weak identity detail.",
                "Started revised identity training; validation LPIPS improves 6.3%.",
                "Benchmarked real GPU batching on a separate 4090.",
            ],
        ),
        dict(
            title="Scope and experimental boundaries",
            bullets=[
                "Target: LUNA; independent code. LHM is a reference, not interchangeable.",
                "User choices: SMPL instead of MHR; NeuMan first.",
                "Six seen subjects; held-out frames. Unseen-person generalization not tested.",
                "Yonsei results are separate from earlier PARCC/B200 work.",
                "No pretrained MV-LHM reconstruction initialization; no larger mixed corpus.",
            ],
        ),
        dict(
            title="Official NeuMan split",
            image="quantitative/charts/dataset-split.png",
            caption="344 train / 44 validation / 41 test. Same identities across splits.",
        ),
        dict(
            title="Implemented pipeline",
            image="quantitative/charts/pipeline.png",
            caption="No canonical RGB target is needed; canonical predictions receive posed-image supervision.",
        ),
        dict(
            title="Verification before long training",
            bullets=[
                "25 CPU tests passed; all 429 real SMPL fits audited.",
                "1,287 original body/face/motion feature tensors verified.",
                "Full-size 4090 smoke and real 512px two-stage resume smoke passed.",
                "Fixed-frame learning pilot improved; strict pixel-replay assertion failed.",
                "Follow-up: exact saved-state restoration, small CUDA repeat variability.",
            ],
        ),
        dict(
            title="Original identity: completed 10,000 updates",
            bullets=[
                "Sapiens body + DINOv2 face; reconstruction weights initialized randomly.",
                "Four references, one target per prediction; effective batch 16.",
                "Validation: LPIPS 0.05796 • PSNR 22.09 dB.",
                "Test: LPIPS 0.05774 • PSNR 21.97 dB.",
                "Faces, hands and fine clothing details remain weak.",
            ],
        ),
        dict(
            title="Original canonical identity output",
            image="qualitative/01_original_identity_all_frames/canonical/citron/views.jpg",
            caption="Virtual camera renders, front / side / back. These are predictions, not canonical GT comparisons.",
        ),
        dict(
            title="Original identity animated by fitted SMPL/LBS",
            image="qualitative/02_original_identity_gt_canonical_lbs_gifs/test/citron-preview.jpg",
            caption="GT | fixed canonical identity | fitted pose. Animated GIF supplied alongside this static slide.",
        ),
        dict(
            title="Neural animator: local articulation failure",
            image="qualitative/13_slide_figures/animator-failure.png",
            caption="Selected animator checkpoint 5k. The LBS and neural-animation columns are different outputs.",
        ),
        dict(
            title="Animator training did not repair the failure",
            image="quantitative/charts/animator-validation.png",
            caption="10k updates completed; best at 5k. Best validation LPIPS 0.2823, test 0.2642.",
        ),
        dict(
            title="Animator diagnosis and small probes",
            bullets=[
                "Local position deltas collapsed across points and driving inputs.",
                "First local-decoder SiLU was saturated; upstream gradients nearly vanished.",
                "FP32 alone did not fix it; normalization probe did not recover articulation.",
                "A fresh local branch could fit two training frames in 512 steps.",
                "This is diagnostic evidence, not a repaired general animator.",
            ],
        ),
        dict(
            title="Identity diagnosis: measured issues",
            bullets=[
                "Extreme attention logits: approximately −1336 to +1628.",
                "Only 9.6–14.6% foreground in inspected square crops.",
                "One target per canonical prediction; original validation still improving at 10k.",
                "Every original update clipped at 0.1; causal effect not established.",
                "Random reconstruction initialization and only six identities remain limitations.",
            ],
        ),
        dict(
            title="Direct fitting reveals remaining representation capacity",
            image="qualitative/13_slide_figures/direct-fit-diagnostic.png",
            caption="128-step direct Gaussian fitting on a training target: LPIPS 0.0492 → 0.00929. Not held-out performance.",
        ),
        dict(
            title="Revised identity training: explicit changes",
            bullets=[
                "Sapiens body and face; per-head Q/K RMS normalization; FP32 decoder.",
                "Four targets share one canonical prediction; effective batch stays 16.",
                "Equal foreground/background RGB weighting.",
                "20k schedule, 250-step warmup, LR 2e−4, clip 1.0.",
                "Several changes together: improvement cannot identify one causal factor.",
            ],
        ),
        dict(
            title="Validation reconstruction improves",
            image="quantitative/charts/identity-validation.png",
            caption=f"Matched 10k: 4.7% lower LPIPS. Selected {new_step:,}: 6.3% lower than completed original identity.",
        ),
        dict(
            title="Same-frame improvement across subjects",
            image="quantitative/charts/per-scene-lpips.png",
            caption="Both frozen models independently rendered on the same 44 validation frames.",
        ),
        dict(
            title="Foreground reconstruction improves",
            image="quantitative/charts/foreground-error.png",
            caption="Selected checkpoint: foreground L1 0.10193 → 0.08881, approximately 12.9% lower.",
        ),
    ]
    for scene in scenes:
        slides.append(
            dict(
                title=f"Qualitative comparison: {scene}",
                image=q + f"{scene}-comparison.png",
                caption=f"GT | original identity 10k | revised identity {new_step:,}. Middle validation frame; fitted SMPL/LBS.",
            )
        )
    slides.extend(
        [
            dict(
                title="Retrained canonical identity and fitted-pose animation",
                image=q + "citron-gt-canonical-lbs.png",
                caption="The middle avatar stays fixed in the supplied GIF. This does not demonstrate a new neural animator.",
            ),
            dict(
                title="Larger GPU batches: measured result",
                image="quantitative/charts/batching.png",
                caption="Two identities fit, but do not speed the encoder. Batched target supervision yields only ~3%.",
            ),
            dict(
                title="Why training remains expensive",
                bullets=[
                    "28,672 tokens per identity reconstruction; five attention blocks.",
                    "FlashAttention already in use.",
                    "Attention forward + backward accounts for ~69% of profiled CUDA operator time.",
                    "Larger encoder batches increase peak allocation: 9.68 → 17.19 GiB.",
                    "The production run was not switched to the experimental batch path.",
                ],
            ),
            dict(
                title="What remains open",
                bullets=[
                    "Complete revised identity training and evaluate the selected test checkpoint.",
                    "Repair/retrain the local neural animator and verify driver-dependent articulation.",
                    "Run controlled ablations and obtain larger reconstruction training data.",
                    "Evaluate unseen identities, cross-identity control and temporal stability.",
                    "Run native released baselines before claiming publication-level comparisons.",
                ],
            ),
            dict(
                title="Conclusion",
                bullets=[
                    "Verified end-to-end development pipeline and reproducible evidence.",
                    "Identity reconstruction improves: 6.3% lower validation LPIPS; 12.9% lower foreground L1.",
                    "Qualitative detail remains imperfect; neural animation remains a failure.",
                    "Larger encoder batches do not materially help throughput.",
                    "Full report, all figures/GIFs, CSVs and audit records are in this folder.",
                ],
            ),
        ]
    )
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def textbox(slide, x, y, w, h, text, size, color="#193048", bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
            p.space_after = Pt(16)
        return shape

    outline = [
        "# Suggested presentation sequence",
        "",
        f"{len(slides)} starter slides; choose the subset suited to the meeting. Text is editable. Figures are static; GIFs are linked in the report and gallery.",
        "",
    ]
    for number, spec in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        textbox(slide, 0.45, 0.25, 12.4, 0.72, spec["title"], 29, bold=True)
        if "image" in spec:
            path = OUT / spec["image"]
            with Image.open(path) as im:
                width, height = im.size
            scale = min(12.3 / width, 5.55 / height)
            w, h = width * scale, height * scale
            slide.shapes.add_picture(
                str(path),
                Inches((13.333 - w) / 2),
                Inches(1.05 + (5.55 - h) / 2),
                width=Inches(w),
                height=Inches(h),
            )
            textbox(slide, 0.5, 6.62, 12.2, 0.52, spec["caption"], 14)
            slide.notes_slide.notes_text_frame.text = f"{spec['caption']}\n\nSource asset: {spec['image']}\nFull protocol and limitations: REPORT.html"
            outline.extend(
                [
                    f"## {number}. {spec['title']}",
                    "",
                    spec["caption"],
                    "",
                    f"Figure: [{Path(spec['image']).name}]({spec['image']})",
                    "",
                ]
            )
        else:
            textbox(
                slide,
                0.7,
                1.5,
                11.9,
                5.15,
                "\n".join("• " + b for b in spec["bullets"]),
                24 if number == 1 else 22,
            )
            slide.notes_slide.notes_text_frame.text = (
                "\n".join(spec["bullets"]) + "\n\nEvidence and qualifications: REPORT.html"
            )
            outline.extend(
                [f"## {number}. {spec['title']}", ""] + ["- " + b for b in spec["bullets"]] + [""]
            )
        textbox(
            slide,
            0.45,
            7.2,
            12,
            0.22,
            f"LUNA-open  |  Yonsei  |  Sept 28, 2026  |  Frozen report snapshot                                      {number}/{len(slides)}",
            10,
            color="#607084",
        )
    prs.save(OUT / "LUNA-open-progress.pptx")
    (OUT / "SLIDE_OUTLINE.md").write_text("\n".join(outline))
    (OUT / "README.txt").write_text(
        "LUNA-open presentation package\n\nOpen REPORT.html for the complete report, gallery.html for all qualitative results, or LUNA-open-progress.pptx for the editable starter slides. Keep the folder structure when copying. Static figures are embedded in the slides; animated GIF files are supplied separately. Source model weights and the licensed SMPL asset are not included. This is a frozen September 28 report, not a live status dashboard.\n"
    )
    summary = dict(
        built_at=datetime.now().astimezone().isoformat(),
        host=socket.gethostname(),
        job_id=os.environ["SLURM_JOB_ID"],
        slides=len(slides),
        media_files=len(media_inventory),
        report_words=len(report.split()),
        models=model_rows,
        same_frame=same_rows,
        old_new_aggregate_metrics_match_saved=True,
        planned_retraining_updates=20000,
        snapshot_training_update=snapshot["training_update_at_snapshot"],
    )
    (OUT / "evidence/build-summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    print(
        json.dumps(
            {k: v for k, v in summary.items() if k not in ("models", "same_frame")}, indent=2
        )
    )


if __name__ == "__main__":
    main()
