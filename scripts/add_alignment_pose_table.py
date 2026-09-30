"""Add the requested editable pose-comparison table and refresh presentation media."""

import argparse
import json
import os
import re
import shutil
import socket
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROWS = [
    ["Method", "Evaluation", "PSNR ↑", "L1 ↓", "LPIPS ↓"],
    ["LHM", "Original", "19.861", "0.02457", "0.09434"],
    ["LHM", "Pose refinement*", "24.895", "0.01364", "0.06740"],
    ["LHM++", "Original", "19.705", "0.02265", "0.07886"],
    ["LHM++", "Pose refinement*", "24.327", "0.01309", "0.05320"],
]
TITLE = "Original vs. pose-refined reconstruction"


def add_pose_comparison_slide(prs, position=10):
    """Insert a native PowerPoint table after the original benchmark slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def text(x, y, width, height, content, size):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = True
        for i, line in enumerate(content.split("\n")):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line
            p.font.size = Pt(size)
            p.font.color.rgb = RGBColor.from_string("243448")
            p.space_after = Pt(10)

    text(.4, .15, 12.5, .8, TITLE, 28)
    text(.6, 1.05, 12.1, .5, "41 NeuMan test frames • 512×512 crops • equal mean over six scenes", 18)
    table = slide.shapes.add_table(5, 5, Inches(.6), Inches(1.85), Inches(12.1), Inches(3.5)).table
    for column, width in zip(table.columns, (1.6, 4.0, 2.2, 2.15, 2.15)):
        column.width = Inches(width)
    for i, row in enumerate(ROWS):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(.16)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                "243448" if i == 0 else "E6F2ED" if i in (2, 4) else "F1F4F8"
            )
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(21)
            p.font.bold = i == 0 or j == 0 or (i in (2, 4) and j >= 2)
            p.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "243448")
            p.alignment = PP_ALIGN.LEFT if j < 2 else PP_ALIGN.RIGHT
    text(.6, 5.8, 12.1, 1.05,
         "* Pose refinement fits root/joint poses to test RGB; identity, shape and camera stay fixed.\n"
         "Diagnostic results; do not substitute for feed-forward benchmark scores.", 17)
    text(.5, 7.08, 12, .32, "Yonsei / fixed 41 test frames / 11 — requested metric comparison", 10)
    ids = prs.slides._sldIdLst
    element = ids[-1]
    ids.remove(element)
    ids.insert(position, element)
    for number, current in enumerate(prs.slides, 1):
        for shape in current.shapes:
            if shape.has_text_frame and shape.text.startswith("Yonsei / fixed 41 test frames /"):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.text = re.sub(r"/ \d{2} —", f"/ {number:02d} —", run.text)
    return slide


def main():
    if not os.environ.get("SLURM_JOB_ID") or "login" in socket.gethostname():
        raise RuntimeError("Slurm compute node required")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    out = args.output.resolve()
    deck = out / "Alignment-investigation.pptx"
    prs = Presentation(deck)
    if any(TITLE in s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame):
        raise RuntimeError("The requested table is already present")
    archive = out.with_suffix(".zip")
    backup = out.parent / "revisions" / f"{out.name}-before-pose-table-{os.environ['SLURM_JOB_ID']}"
    backup.mkdir(parents=True, exist_ok=False)
    for source in (deck, archive, archive.with_suffix(".receipt.json"), out / "AUDIT.json", out / "BUILD.json", out / "SHA256.json", out / "Alignment-investigation-preview.pdf"):
        shutil.copy2(source, backup / source.name)
    add_pose_comparison_slide(prs)
    prs.save(deck)
    from audit_neuman_final_report import digest, slide_previews

    slide_previews(prs, out)
    (out / "NeuMan-final-report-preview.pdf").replace(out / "Alignment-investigation-preview.pdf")
    revision = dict(host=socket.gethostname(), job=os.environ["SLURM_JOB_ID"], slide=11,
                    slides=len(prs.slides), rows=ROWS, backup=str(backup),
                    change="Added editable original versus pose-refinement table; metrics unchanged",
                    previews="Regenerated using Pillow; Office playback not executed")
    (out / "PRESENTATION-UPDATE.json").write_text(json.dumps(revision, indent=2))
    for filename in ("BUILD.json", "AUDIT.json"):
        path = out / filename
        record = json.loads(path.read_text())
        record["slides"] = len(prs.slides)
        record["presentation_revision"] = revision
        path.write_text(json.dumps(record, indent=2))
    for filename in ("add_alignment_pose_table.py", "build_alignment_report.py"):
        shutil.copy2(Path(__file__).parent / filename, out / "evidence/scripts" / filename)
    hashes = {str(p.relative_to(out)): digest(p) for p in sorted(out.rglob("*")) if p.is_file() and p.name != "SHA256.json"}
    (out / "SHA256.json").write_text(json.dumps(hashes, indent=2))
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=5) as z:
        for path in sorted(out.rglob("*")):
            if path.is_file():
                z.write(path, path.relative_to(out.parent))
    receipt = dict(path=str(archive), bytes=archive.stat().st_size, sha256=digest(archive),
                   files=len(hashes)+1, presentation_revision=revision,
                   audit=json.loads((out / "AUDIT.json").read_text()))
    archive.with_suffix(".receipt.json").write_text(json.dumps(receipt, indent=2))
    print(json.dumps({k: v for k, v in receipt.items() if k != "audit"}, indent=2), flush=True)


if __name__ == "__main__":
    main()
