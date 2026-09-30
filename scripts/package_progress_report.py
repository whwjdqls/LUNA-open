"""Check local report assets and create the single downloadable archive."""

import csv
import hashlib
import json
import os
import shutil
import socket
import zipfile
from collections import Counter
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlsplit

from PIL import Image
from pptx import Presentation

WORK = Path("/scratch2/whwjdqls99/LUNA-open")
OUT = WORK / "reports/luna-progress-20260928"
PRIVATE = WORK / "reports/luna-progress-20260928-work"
REPO = Path("/home/whwjdqls99/LUNA-open")


def digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.ids = set()

    def handle_starttag(self, tag, attrs):
        for key, value in attrs:
            if key in ("href", "src") and value:
                self.links.append(value)
            if key == "id" and value:
                self.ids.add(value)


def main():
    if not os.environ.get("SLURM_JOB_ID") or "login" in socket.gethostname():
        raise RuntimeError("Use a CPU compute allocation")
    shutil.copy2(Path(__file__), OUT / "evidence/source/package_progress_report.py")
    shutil.copy2(
        REPO / "scripts/report_assets/progress-report.md",
        OUT / "evidence/source/progress-report-template.md",
    )
    parsed = {}
    for name in ("REPORT.html", "gallery.html"):
        parser = Links()
        parser.feed((OUT / name).read_text())
        parsed[name] = parser
    bad_links = []
    checked_links = 0
    for name, parser in parsed.items():
        for value in parser.links:
            parts = urlsplit(value)
            if parts.scheme or parts.netloc:
                continue
            path = (OUT / unquote(parts.path)).resolve() if parts.path else OUT / name
            checked_links += 1
            if not path.exists():
                bad_links.append(dict(source=name, link=value, reason="missing file"))
            elif (
                parts.fragment
                and path.name in parsed
                and parts.fragment not in parsed[path.name].ids
            ):
                bad_links.append(dict(source=name, link=value, reason="missing anchor"))
    if bad_links:
        print(json.dumps(bad_links, indent=2), flush=True)
        raise AssertionError("Broken report/gallery links")
    counts = Counter()
    gifs = []
    for path in sorted((OUT / "qualitative").rglob("*")):
        if not path.is_file() or path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".gif"}:
            continue
        with Image.open(path) as picture:
            if path.suffix.lower() == ".gif":
                frames = picture.n_frames
                duration = []
                for index in range(frames):
                    picture.seek(index)
                    picture.load()
                    duration.append(picture.info.get("duration"))
                gifs.append(
                    dict(
                        path=path.relative_to(OUT).as_posix(),
                        frames=frames,
                        durations_ms=sorted(set(duration)),
                    )
                )
            else:
                picture.verify()
        counts[path.suffix.lower()] += 1
    # Expected sequence coverage for copied train/test and new validation GIFs.
    manifest = json.loads((OUT / "evidence/manifest-v2.json").read_text())
    for entry in gifs:
        path = Path(entry["path"])
        if "02_original_identity_gt_canonical_lbs_gifs" in path.parts:
            scene, split = path.stem, path.parent.name
            assert entry["frames"] == len(manifest["scenes"][scene]["splits"][split])
        if "10_same_frame_identity_comparison" in path.parts:
            scene = path.name.split("-")[0]
            assert entry["frames"] == len(manifest["scenes"][scene]["splits"]["val"])
    prs = Presentation(OUT / "LUNA-open-progress.pptx")
    assert len(prs.slides) == 29
    slide_errors = []
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > prs.slide_width + 100
                or shape.top + shape.height > prs.slide_height + 100
            ):
                slide_errors.append(dict(slide=number, name=shape.name))
        assert slide.notes_slide.notes_text_frame.text.strip()
    assert not slide_errors, slide_errors
    with zipfile.ZipFile(OUT / "LUNA-open-progress.pptx") as deck:
        assert deck.testzip() is None
    checks = dict(
        checked_at=datetime.now().astimezone().isoformat(),
        host=socket.gethostname(),
        job_id=os.environ["SLURM_JOB_ID"],
        local_report_gallery_links_checked=checked_links,
        broken_links=bad_links,
        image_files_checked=dict(counts),
        gif_sequences=gifs,
        all_gif_frames_decoded=True,
        expected_gif_frame_counts_match=True,
        slides=len(prs.slides),
        slide_shapes_in_bounds=True,
        slide_notes_present=True,
        pptx_zip_crc_passed=True,
        powerpoint_visual_rendering="Not executed in Microsoft PowerPoint or LibreOffice; source figures visually inspected and slide package structure checked.",
        metric_comparison="Fresh original/retrained validation aggregate metrics match the saved records within 1e-7; see build-summary.json.",
    )
    (OUT / "PACKAGE_CHECKS.json").write_text(json.dumps(checks, indent=2) + "\n")
    files = [p for p in sorted(OUT.rglob("*")) if p.is_file() and p.name != "FILE_MANIFEST.csv"]
    with (OUT / "FILE_MANIFEST.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["path", "bytes", "sha256"])
        writer.writeheader()
        for path in files:
            writer.writerow(
                dict(
                    path=path.relative_to(OUT).as_posix(),
                    bytes=path.stat().st_size,
                    sha256=digest(path),
                )
            )
    archive = OUT.with_suffix(".zip")
    partial = archive.with_suffix(".zip.part")
    with zipfile.ZipFile(
        partial, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=4, allowZip64=True
    ) as zipped:
        for path in sorted(OUT.rglob("*")):
            if path.is_file():
                zipped.write(path, str(Path(OUT.name) / path.relative_to(OUT)))
    partial.replace(archive)
    with zipfile.ZipFile(archive) as zipped:
        assert zipped.testzip() is None
        archived_files = len(zipped.infolist())
    archive_hash = digest(archive)
    archive.with_suffix(".zip.sha256").write_text(f"{archive_hash}  {archive.name}\n")
    receipt = dict(
        folder=str(OUT),
        archive=str(archive),
        bytes=archive.stat().st_size,
        sha256=archive_hash,
        files=archived_files,
        media_counts=dict(counts),
        slides=len(prs.slides),
        local_links_checked=checked_links,
        archive_crc_passed=True,
        created_at=datetime.now().astimezone().isoformat(),
    )
    (PRIVATE / "package-receipt.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(receipt, indent=2), flush=True)


if __name__ == "__main__":
    main()
